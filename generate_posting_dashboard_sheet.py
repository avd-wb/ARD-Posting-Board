#!/usr/bin/env python3
"""
generate_posting_dashboard_sheet.py
Generates the comprehensive, dynamic Google Sheets-ready workbook:
WB_ARD_Interactive_Posting_Board_GoogleSheets_Ready.xlsx

Key Features & Enhancements:
1. Full Officer Names: Verified complete names (e.g. Dr. Soma Das (nee Saha) for HRMS 1995004945).
2. Dynamic Dropdown Option Reduction across BOTH Substantive (Col J) and Service Utilized (Col L) columns:
   - A post will ONLY NOT SHOW in dropdowns if it has been used up somewhere in that column OR in the service utilized column!
   - Dynamic FILTER arrays in Available_DD_Posts, Available_AD_Vacancies, and All_Cadre_Posts_1794.
3. Cross-Column Duplicate Detection & Prominent Red Warning:
   - Selecting a duplicate post in Column J or Column L highlights the cell in BRIGHT RED (#FEE2E2 / #991B1B).
   - Real-time status flags duplicate across both substantive and SU columns.
   - Column N displays actionable error.
   - Master KPI card counts duplicate allotments in red.
   - Master availability sheets flag duplicate with overallocated status.
4. Block-Level Post Nomenclature:
   - Block posts (ABAHC, BAHC, BLDO, SAHC, VO) explicitly display the block name:
     e.g., "[291] Veterinary Officer, BAHC, Darjeeling Pulbazar Block (Darjeeling)"
5. Complete Cadre, Order Schedule & Legend Sheets.
"""

import os
import sqlite3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

DB_PATH = "ard_master_truth.db"
OUTPUT_FILE = "WB_ARD_Interactive_Posting_Board_GoogleSheets_Ready.xlsx"

def format_cadre_label(psl, desig, estab, block, dist):
    """
    Formats post labels to always prominently display the Block name
    for block-level posts (ABAHC, BAHC, BLDO, SAHC, VO).
    """
    dist = dist or "HQ"
    estab = (estab or "").strip()
    block = (block or "").strip()
    
    is_generic_setup = any(gen in estab.lower() for gen in [
        'block level set up', 'block/sub-divisional level set up', 
        'sub-divisional and block level set up', 'district set up'
    ])
    is_block_post = any(kw in desig.upper() for kw in ['ABAHC', 'BAHC', 'BLDO', 'SAHC', 'BLOCK LIVESTOCK', 'VO,'])
    
    if is_block_post or is_generic_setup:
        if block and block.lower() not in ['-', 'none', '']:
            b_str = block if ('block' in block.lower() or 'hq' in block.lower()) else f'{block} Block'
            return f"[{psl}] {desig}, {b_str} ({dist})"
        elif estab:
            return f"[{psl}] {desig} ({estab}) - {dist}"
        else:
            return f"[{psl}] {desig} - {dist}"
            
    is_hq_block = any(h in block.lower() for h in ['directorate hq', 'directorate headquarter', 'hq', 'headquarter', '-'])
    if estab and block and not is_hq_block and block.lower() not in estab.lower() and block.lower() not in dist.lower():
        b_str = block if 'block' in block.lower() else f'{block} Block'
        return f"[{psl}] {desig} ({estab}, {b_str}) - {dist}"
    elif estab:
        return f"[{psl}] {desig} ({estab}) - {dist}"
    else:
        return f"[{psl}] {desig} - {dist}"

def build_interactive_dashboard():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    c = conn.cursor()

    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Remove default sheet

    # Colors & Fonts
    c_navy = "1E3A8A"
    c_teal = "0F766E"
    c_indigo = "3730A3"
    c_amber = "92400E"
    c_slate = "334155"

    fill_navy = PatternFill(start_color=c_navy, end_color=c_navy, fill_type="solid")
    fill_teal = PatternFill(start_color=c_teal, end_color=c_teal, fill_type="solid")
    fill_indigo = PatternFill(start_color=c_indigo, end_color=c_indigo, fill_type="solid")
    fill_amber = PatternFill(start_color=c_amber, end_color=c_amber, fill_type="solid")
    fill_slate = PatternFill(start_color=c_slate, end_color=c_slate, fill_type="solid")

    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_plain = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    font_white_hdr = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_sec_hdr = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    font_title = Font(name="Calibri", size=15, bold=True, color=c_navy)
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_normal = Font(name="Calibri", size=10)
    font_italic = Font(name="Calibri", size=10, italic=True, color="64748B")
    font_kpi_num = Font(name="Calibri", size=18, bold=True, color=c_navy)
    font_kpi_lbl = Font(name="Calibri", size=9, bold=True, color="475569")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # Red warning styles
    fill_red_warning = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    font_red_bold = Font(name="Calibri", size=10, bold=True, color="991B1B")

    # =========================================================================
    # TAB 3: Available_DD_Posts (Create first so formulas can reference it)
    # =========================================================================
    ws_dd = wb.create_sheet(title="Available_DD_Posts")
    ws_dd.views.sheetView[0].showGridLines = True

    dd_headers = [
        "Post ID", "Post Name & Station (Master Identifier)", "Establishment", 
        "District", "Total Times Selected (Col J + Col L)", "Live Availability Status", 
        "Allotted Officer Name", "Allocation Alert & Notes", 
        "Remaining Available Posts (Dynamic Dropdown Source for Column J)"
    ]
    for c_idx, h in enumerate(dd_headers, 1):
        cell = ws_dd.cell(row=1, column=c_idx, value=h)
        cell.fill = fill_navy if c_idx != 9 else fill_indigo
        cell.font = font_white_hdr
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
    ws_dd.row_dimensions[1].height = 28

    c.execute("SELECT dd_sl, post_name, establishment, district FROM available_dd_posts WHERE is_blocked_vigilance = 0 ORDER BY dd_sl")
    dd_rows = c.fetchall()
    for idx, r in enumerate(dd_rows, 2):
        post_id = r["dd_sl"]
        dist = r["district"] or "HQ"
        estab = r["establishment"] or ""
        pname = r["post_name"] or "Deputy Director, ARD"
        label = f"[DD-{post_id}] {pname} ({estab}) - {dist}"

        ws_dd.cell(row=idx, column=1, value=post_id).alignment = Alignment(horizontal="center")
        ws_dd.cell(row=idx, column=2, value=label).font = font_bold
        ws_dd.cell(row=idx, column=3, value=estab)
        ws_dd.cell(row=idx, column=4, value=dist).alignment = Alignment(horizontal="center")
        
        # Total selections from BOTH Substantive Col J AND Service Utilized Col L
        ws_dd.cell(row=idx, column=5, value=f"=COUNTIF(Posting_Dashboard!$J$10:$J$360, B{idx}) + COUNTIF(Posting_Dashboard!$L$10:$L$360, B{idx})").alignment = Alignment(horizontal="center")
        # Status Formula
        ws_dd.cell(row=idx, column=6, value=f'=IF(E{idx}=0, "AVAILABLE", IF(E{idx}=1, "ALLOTTED", "⚠️ DUPLICATE ALLOTMENT!"))').alignment = Alignment(horizontal="center")
        # Allotted Officer formula (checks both Col J and Col L)
        ws_dd.cell(row=idx, column=7, value=f'=IF(E{idx}>0, IFERROR(INDEX(Posting_Dashboard!$D$10:$D$360, MATCH(B{idx}, Posting_Dashboard!$J$10:$J$360, 0)), IFERROR(INDEX(Posting_Dashboard!$D$10:$D$360, MATCH(B{idx}, Posting_Dashboard!$L$10:$L$360, 0)), "-")), "-")')
        # Allocation Alert
        ws_dd.cell(row=idx, column=8, value=f'=IF(E{idx}>1, "⚠️ OVERALLOCATED (" & E{idx} & " times)", IF(E{idx}=1, "Allotted to " & G{idx}, "Available for Selection"))')

        fill_color = fill_zebra if idx % 2 == 0 else fill_plain
        for c_idx in range(1, 9):
            cell = ws_dd.cell(row=idx, column=c_idx)
            cell.border = cell_border
            if c_idx not in [2, 6, 8]:
                cell.fill = fill_color
            cell.font = font_normal if c_idx != 2 else font_bold
        ws_dd.row_dimensions[idx].height = 20

    # Column 9: Dynamic Filter Formula for Dropdowns - Only shows posts NOT used in Col J or Col L
    ws_dd.cell(row=2, column=9, value=f'=IFERROR(FILTER(B$2:B${len(dd_rows)+1}, E$2:E${len(dd_rows)+1}=0), "ALL POSTS ALLOTTED")').font = font_bold

    # =========================================================================
    # TAB 4: Available_AD_Vacancies
    # =========================================================================
    ws_ad = wb.create_sheet(title="Available_AD_Vacancies")
    ws_ad.views.sheetView[0].showGridLines = True

    ad_headers = [
        "Post ID", "Post Name & Station (Dropdown Identifier)", "Establishment", 
        "District", "Total Times Selected (Col J + Col L)", "Live Availability Status", 
        "Allotted Officer Name", "Allocation Alert & Notes",
        "Remaining Available AD Vacancies (Dynamic Dropdown Source for Column J)"
    ]
    for c_idx, h in enumerate(ad_headers, 1):
        cell = ws_ad.cell(row=1, column=c_idx, value=h)
        cell.fill = fill_teal if c_idx != 9 else fill_indigo
        cell.font = font_white_hdr
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
    ws_ad.row_dimensions[1].height = 28

    c.execute("""
        SELECT id, post_sl, designation, establishment, block, district 
        FROM cadre_1794_posts 
        WHERE occupancy_status = 'Vacant' AND designation LIKE '%Assistant Director%'
        ORDER BY post_sl
    """)
    ad_rows = c.fetchall()
    for idx, r in enumerate(ad_rows, 2):
        psl = r["post_sl"]
        desig = r["designation"]
        estab = r["establishment"] or ""
        block = r["block"] or ""
        dist = r["district"] or "HQ"
        label = format_cadre_label(psl, desig, estab, block, dist)

        ws_ad.cell(row=idx, column=1, value=psl).alignment = Alignment(horizontal="center")
        ws_ad.cell(row=idx, column=2, value=label).font = font_bold
        ws_ad.cell(row=idx, column=3, value=estab)
        ws_ad.cell(row=idx, column=4, value=dist).alignment = Alignment(horizontal="center")
        # Total selections from BOTH Substantive Col J AND Service Utilized Col L
        ws_ad.cell(row=idx, column=5, value=f"=COUNTIF(Posting_Dashboard!$J$10:$J$360, B{idx}) + COUNTIF(Posting_Dashboard!$L$10:$L$360, B{idx})").alignment = Alignment(horizontal="center")
        ws_ad.cell(row=idx, column=6, value=f'=IF(E{idx}=0, "AVAILABLE", IF(E{idx}=1, "ALLOTTED", "⚠️ DUPLICATE!"))').alignment = Alignment(horizontal="center")
        ws_ad.cell(row=idx, column=7, value=f'=IF(E{idx}>0, IFERROR(INDEX(Posting_Dashboard!$D$10:$D$360, MATCH(B{idx}, Posting_Dashboard!$J$10:$J$360, 0)), IFERROR(INDEX(Posting_Dashboard!$D$10:$D$360, MATCH(B{idx}, Posting_Dashboard!$L$10:$L$360, 0)), "-")), "-")')
        ws_ad.cell(row=idx, column=8, value=f'=IF(E{idx}>1, "⚠️ OVERALLOCATED (" & E{idx} & " times)", IF(E{idx}=1, "Allotted to " & G{idx}, "Available for Selection"))')

        fill_color = fill_zebra if idx % 2 == 0 else fill_plain
        for c_idx in range(1, 9):
            cell = ws_ad.cell(row=idx, column=c_idx)
            cell.border = cell_border
            if c_idx not in [2, 6, 8]:
                cell.fill = fill_color
            cell.font = font_normal if c_idx != 2 else font_bold
        ws_ad.row_dimensions[idx].height = 20

    # Column 9: Dynamic Filter Formula for AD Vacancy Dropdowns
    ws_ad.cell(row=2, column=9, value=f'=IFERROR(FILTER(B$2:B${len(ad_rows)+1}, E$2:E${len(ad_rows)+1}=0), "ALL VACANCIES ALLOTTED")').font = font_bold

    # =========================================================================
    # TAB 5: All_Cadre_Posts_1794 (Complete SU lookup source + DD Posts)
    # =========================================================================
    ws_su = wb.create_sheet(title="All_Cadre_Posts_1794")
    ws_su.views.sheetView[0].showGridLines = True

    cadre_headers = [
        "Post Label (Dropdown Identifier)", "Designation & Office", "Occupancy Status", 
        "Current Incumbent & HRMS", "District", "Pay Level", "Post ID",
        "Total Times Selected (Col L + Col J)", "Live Utilization Status",
        "Remaining Available Posts for Service Utilization (Dynamic Dropdown Source for Column L)"
    ]
    for c_idx, h in enumerate(cadre_headers, 1):
        cell = ws_su.cell(row=1, column=c_idx, value=h)
        cell.fill = fill_slate if c_idx != 10 else fill_indigo
        cell.font = font_white_hdr
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
    ws_su.row_dimensions[1].height = 28

    c.execute("""
        SELECT post_sl, designation, establishment, block, district, occupancy_status, incumbent_name, incumbent_hrms, pay_level 
        FROM cadre_1794_posts 
        ORDER BY post_sl
    """)
    all_cadre = c.fetchall()
    
    # 1. Populate all 1,794 Cadre Posts (Rows 2 to 1795)
    for idx, r in enumerate(all_cadre, 2):
        psl = r["post_sl"]
        desig = r["designation"]
        estab = r["establishment"] or ""
        block = r["block"] or ""
        dist = r["district"] or "HQ"
        occ = r["occupancy_status"] or "Vacant"
        inc_name = r["incumbent_name"] or "None"
        inc_hrms = r["incumbent_hrms"] or ""
        inc_str = f"{inc_name} ({inc_hrms})" if inc_hrms else ("Vacant" if occ == "Vacant" else inc_name)
        label = format_cadre_label(psl, desig, estab, block, dist)

        is_block_post = any(kw in desig.upper() for kw in ['ABAHC', 'BAHC', 'BLDO', 'SAHC', 'BLOCK LIVESTOCK'])
        if is_block_post and block:
            b_str = block if 'block' in block.lower() or 'hq' in block.lower() else f'{block} Block'
            office_desc = f"{desig}, {b_str} ({dist})"
        else:
            office_desc = f"{desig} - {estab}" if estab else desig

        ws_su.cell(row=idx, column=1, value=label).font = font_bold
        ws_su.cell(row=idx, column=2, value=office_desc)
        ws_su.cell(row=idx, column=3, value=occ).alignment = Alignment(horizontal="center")
        ws_su.cell(row=idx, column=4, value=inc_str)
        ws_su.cell(row=idx, column=5, value=dist).alignment = Alignment(horizontal="center")
        ws_su.cell(row=idx, column=6, value=r["pay_level"]).alignment = Alignment(horizontal="center")
        ws_su.cell(row=idx, column=7, value=psl).alignment = Alignment(horizontal="center")
        
        # Col 8 (H): Total selections in SU (Col L) + Substantive (Col J)
        ws_su.cell(row=idx, column=8, value=f"=COUNTIF(Posting_Dashboard!$L$10:$L$360, A{idx}) + COUNTIF(Posting_Dashboard!$J$10:$J$360, A{idx})").alignment = Alignment(horizontal="center")
        # Col 9 (I): Live utilization status
        ws_su.cell(row=idx, column=9, value=f'=IF(H{idx}=0, C{idx}, IF(H{idx}=1, "UTILIZED IN DASHBOARD", "⚠️ DUPLICATE ALLOTMENT!"))').alignment = Alignment(horizontal="center")

        fill_color = fill_zebra if idx % 2 == 0 else fill_plain
        for c_idx in range(1, 10):
            cell = ws_su.cell(row=idx, column=c_idx)
            cell.border = cell_border
            if c_idx not in [3, 9]:
                cell.fill = fill_color
            cell.font = font_normal if c_idx != 1 else font_bold
        ws_su.row_dimensions[idx].height = 19

    # 2. Append the 242 DD posts from Available_DD_Posts so they can ALSO be chosen for SU (Rows 1796 to 2037)
    curr_su_row = len(all_cadre) + 2
    for r in dd_rows:
        post_id = r["dd_sl"]
        dist = r["district"] or "HQ"
        estab = r["establishment"] or ""
        pname = r["post_name"] or "Deputy Director, ARD"
        label = f"[DD-{post_id}] {pname} ({estab}) - {dist}"

        ws_su.cell(row=curr_su_row, column=1, value=label).font = font_bold
        ws_su.cell(row=curr_su_row, column=2, value=f"{pname} - {estab}")
        ws_su.cell(row=curr_su_row, column=3, value="Vacant").alignment = Alignment(horizontal="center")
        ws_su.cell(row=curr_su_row, column=4, value="None (Available for Promotion / SU)")
        ws_su.cell(row=curr_su_row, column=5, value=dist).alignment = Alignment(horizontal="center")
        ws_su.cell(row=curr_su_row, column=6, value="Level-19").alignment = Alignment(horizontal="center")
        ws_su.cell(row=curr_su_row, column=7, value=f"DD-{post_id}").alignment = Alignment(horizontal="center")
        
        ws_su.cell(row=curr_su_row, column=8, value=f"=COUNTIF(Posting_Dashboard!$L$10:$L$360, A{curr_su_row}) + COUNTIF(Posting_Dashboard!$J$10:$J$360, A{curr_su_row})").alignment = Alignment(horizontal="center")
        ws_su.cell(row=curr_su_row, column=9, value=f'=IF(H{curr_su_row}=0, "Vacant", IF(H{curr_su_row}=1, "UTILIZED IN DASHBOARD", "⚠️ DUPLICATE ALLOTMENT!"))').alignment = Alignment(horizontal="center")

        fill_color = fill_zebra if curr_su_row % 2 == 0 else fill_plain
        for c_idx in range(1, 10):
            cell = ws_su.cell(row=curr_su_row, column=c_idx)
            cell.border = cell_border
            if c_idx not in [3, 9]:
                cell.fill = fill_color
            cell.font = font_normal if c_idx != 1 else font_bold
        ws_su.row_dimensions[curr_su_row].height = 19
        curr_su_row += 1

    total_su_rows = curr_su_row - 1

    # Col 10 (J): Dynamic Filter Formula for Service Utilization Dropdown (Excludes any post used in Col L or Col J)
    ws_su.cell(row=2, column=10, value=f'=IFERROR(FILTER(A$2:A${total_su_rows}, H$2:H${total_su_rows}=0), "ALL POSTS UTILIZED")').font = font_bold

    # =========================================================================
    # TAB 1: Posting_Dashboard (THE MASTER TAB - Put at Index 0)
    # =========================================================================
    ws_dash = wb.create_sheet(title="Posting_Dashboard", index=0)
    ws_dash.views.sheetView[0].showGridLines = True

    # Header Masthead
    ws_dash.merge_cells("A1:Q1")
    ws_dash.merge_cells("A2:Q2")
    ws_dash.merge_cells("A3:Q3")

    ws_dash.cell(row=1, column=1, value="GOVERNMENT OF WEST BENGAL — ANIMAL RESOURCES DEVELOPMENT DEPARTMENT").font = Font(name="Calibri", size=15, bold=True, color="1E3A8A")
    ws_dash.cell(row=2, column=1, value="OFFICIAL POSTING & TRANSFER DECISION BOARD (DYNAMIC INTERACTIVE COCKPIT)").font = Font(name="Calibri", size=12, bold=True, color="0F766E")
    ws_dash.cell(row=3, column=1, value="Real-Time Cadre Placement, 50-Point Roster Promotions, Service Utilization & Displacement Cascades").font = font_italic

    for r_idx in range(1, 4):
        ws_dash.cell(row=r_idx, column=1).alignment = Alignment(horizontal="center", vertical="center")
        ws_dash.row_dimensions[r_idx].height = 22

    # Executive Live KPI Cards (Row 5 & 6) - 7 Cards across Columns B to O
    kpi_configs = [
        (2, "TOTAL CADRE POSTS", "1,794", "Sanctioned Posts"),
        (4, "CLEAR VACANCIES", "750", "Available in Cadre"),
        (6, "AVAILABLE DD POSTS", '=COUNTIF(Available_DD_Posts!$F$2:$F$243, "AVAILABLE")', "Deputy Director Seats"),
        (8, "TOTAL ALLOTTED", '=COUNTIF($M$10:$M$360, "ALLOTTED*") + COUNTIF($M$10:$M$360, "REHABILITATED*") + COUNTIF($M$10:$M$360, "REALLOCATED")', "Decisions Made"),
        (10, "DISPLACED DUE TO SU", '=COUNTIF($M$10:$M$360, "*COLLISION*")', "Incumbents Displaced"),
        (12, "PENDING SELECTION", '=COUNTIF($M$10:$M$360, "PENDING*")', "Awaiting Allocation"),
        (14, "⚠️ DUPLICATE ALERTS", '=COUNTIF($M$10:$M$360, "*DUPLICATE*")', "Overallocated Posts"),
    ]

    fill_kpi = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    for c_start, label, val_formula, subtext in kpi_configs:
        ws_dash.merge_cells(start_row=5, start_column=c_start, end_row=5, end_column=c_start+1)
        ws_dash.merge_cells(start_row=6, start_column=c_start, end_row=6, end_column=c_start+1)

        val_cell = ws_dash.cell(row=5, column=c_start, value=val_formula)
        val_cell.font = font_kpi_num
        val_cell.alignment = Alignment(horizontal="center", vertical="center")
        val_cell.fill = fill_kpi
        val_cell.border = cell_border

        lbl_cell = ws_dash.cell(row=6, column=c_start, value=f"{label} ({subtext})")
        lbl_cell.font = font_kpi_lbl
        lbl_cell.alignment = Alignment(horizontal="center", vertical="center")
        lbl_cell.fill = fill_kpi
        lbl_cell.border = cell_border

        ws_dash.cell(row=5, column=c_start+1).border = cell_border
        ws_dash.cell(row=6, column=c_start+1).border = cell_border

    ws_dash.row_dimensions[5].height = 28
    ws_dash.row_dimensions[6].height = 18
    ws_dash.row_dimensions[7].height = 10 # Spacing

    # Section 1 Header (Row 8)
    ws_dash.merge_cells("A8:Q8")
    sec1_cell = ws_dash.cell(row=8, column=1, value="SECTION 1: 50-POINT ROSTER PROMOTION & TRANSFER CONSOLE (242 CANDIDATES TO DEPUTY DIRECTOR)")
    sec1_cell.fill = fill_navy
    sec1_cell.font = font_sec_hdr
    sec1_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_dash.row_dimensions[8].height = 28

    # Table Column Headers (Row 9)
    dashboard_headers = [
        "Sl", "Roster Pt & Quota", "HRMS ID", "Officer Name", 
        "Present Designation & Post", "Present Block", "Present District", "Superannuation (DOR)",
        "Stated Preferences (Choices 1-3)",
        "Eligible Post Picker: Deputy Director (Substantive Dropdown)",
        "Enable Service Utilization? (NO / YES)",
        "Service Utilization (SU) Post Picker (Cadre Dropdown)",
        "Allotment & Transfer Status",
        "Displaced Incumbent Status (Eligible for Transfer due to Displacement)",
        "Statutory Category for Order",
        "Present Pay Level",
        "New Pay Level"
    ]

    for c_idx, h in enumerate(dashboard_headers, 1):
        cell = ws_dash.cell(row=9, column=c_idx, value=h)
        if c_idx in [10, 11, 12]:
            cell.fill = fill_indigo # Action Pickers in Indigo
        elif c_idx in [13, 14]:
            cell.fill = fill_amber # Alerts in Amber
        else:
            cell.fill = fill_navy
        cell.font = font_white_hdr
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
    ws_dash.row_dimensions[9].height = 36

    # Populate 242 Roster Candidates (Rows 10 to 251)
    c.execute("""
        SELECT sl_no, roster_point, point_reserved_for, hrms_id, officer_name, 
               present_posting, present_block, present_district, service_ends, 
               pref_1, pref_2, pref_3,
               substantive_post_name, su_post_name, allotment_status
        FROM roster_50_point_candidates 
        ORDER BY sl_no
    """)
    roster_rows = c.fetchall()

    for idx, r in enumerate(roster_rows, 10):
        ws_dash.cell(row=idx, column=1, value=r["sl_no"]).alignment = Alignment(horizontal="center")
        ws_dash.cell(row=idx, column=2, value=f"Pt {r['roster_point']} ({r['point_reserved_for']})").alignment = Alignment(horizontal="center")
        ws_dash.cell(row=idx, column=3, value=r["hrms_id"]).alignment = Alignment(horizontal="center")
        ws_dash.cell(row=idx, column=4, value=r["officer_name"]).font = font_bold
        ws_dash.cell(row=idx, column=5, value=r["present_posting"])
        ws_dash.cell(row=idx, column=6, value=r["present_block"] or "-")
        ws_dash.cell(row=idx, column=7, value=r["present_district"] or "-")
        ws_dash.cell(row=idx, column=8, value=r["service_ends"] or "-").alignment = Alignment(horizontal="center")

        prefs_str = ", ".join(filter(None, [r["pref_1"], r["pref_2"], r["pref_3"]])) or "No preference recorded"
        ws_dash.cell(row=idx, column=9, value=prefs_str).font = font_italic

        # Dropdown Picker for DD Post (Col J)
        ws_dash.cell(row=idx, column=10, value="").font = font_bold
        # Toggle SU (Col K)
        ws_dash.cell(row=idx, column=11, value="NO").alignment = Alignment(horizontal="center")
        # SU Post Picker (Col L)
        ws_dash.cell(row=idx, column=12, value="").font = font_bold

        # Live Status Formula (Col M) - WITH CROSS-COLUMN DUPLICATE DETECTION (J & L)
        status_formula = (
            f'=IF(ISBLANK(J{idx}), "PENDING", '
            f'IF(OR(COUNTIF($J$10:$J$360, J{idx})>1, COUNTIF($L$10:$L$360, J{idx})>0), '
            f'"⚠️ DUPLICATE ALLOTMENT (" & (COUNTIF($J$10:$J$360, J{idx}) + COUNTIF($L$10:$L$360, J{idx})) & " OFFICERS SELECTED)", '
            f'IF(K{idx}="YES", '
            f'IF(OR(COUNTIF($L$10:$L$360, L{idx})>1, COUNTIF($J$10:$J$360, L{idx})>0), '
            f'"⚠️ DUPLICATE SU ALLOTMENT (" & (COUNTIF($L$10:$L$360, L{idx}) + COUNTIF($J$10:$J$360, L{idx})) & " OFFICERS SELECTED)", '
            f'IF(IFERROR(VLOOKUP(L{idx}, All_Cadre_Posts_1794!$A$2:$C$2038, 3, FALSE), "")="Occupied", "⚠️ SU COLLISION (DISPLACEMENT TRIGGERED)", "ALLOTTED WITH SU")), '
            f'"ALLOTTED DIRECT")))'
        )
        ws_dash.cell(row=idx, column=13, value=status_formula).alignment = Alignment(horizontal="center")

        # Live Displaced Incumbent Formula (Col N) - WITH CROSS-COLUMN DUPLICATE ERROR
        displaced_formula = (
            f'=IF(ISBLANK(J{idx}), "Awaiting Post Selection", '
            f'IF(COUNTIF($J$10:$J$360, J{idx}) + COUNTIF($L$10:$L$360, J{idx}) > 1, '
            f'"⚠️ ERROR: Post duplicate allotted across substantive/SU columns! Reassign post.", '
            f'IF(AND(K{idx}="YES", NOT(ISBLANK(L{idx}))), '
            f'IF(COUNTIF($L$10:$L$360, L{idx}) + COUNTIF($J$10:$J$360, L{idx}) > 1, '
            f'"⚠️ ERROR: SU post duplicate allotted across substantive/SU columns! Reassign SU post.", '
            f'IF(IFERROR(VLOOKUP(L{idx}, All_Cadre_Posts_1794!$A$2:$D$2038, 3, FALSE), "")="Occupied", '
            f'"⚠️ " & IFERROR(VLOOKUP(L{idx}, All_Cadre_Posts_1794!$A$2:$D$2038, 4, FALSE), "Incumbent Doctor") & " — MARKED AS ELIGIBLE OFFICER FOR TRANSFER DUE TO DISPLACEMENT", '
            f'"Clear Vacancy / No Displacement")), '
            f'"Clean Direct Placement")))'
        )
        ws_dash.cell(row=idx, column=14, value=displaced_formula)

        ws_dash.cell(row=idx, column=15, value="Promotion under WBS (ROPA) 2019")
        ws_dash.cell(row=idx, column=16, value="Level 16").alignment = Alignment(horizontal="center")
        ws_dash.cell(row=idx, column=17, value="Level 19 (Rs. 95,100 - Rs. 1,48,000)").alignment = Alignment(horizontal="center")

        fill_color = fill_zebra if idx % 2 == 0 else fill_plain
        for c_idx in range(1, 18):
            cell = ws_dash.cell(row=idx, column=c_idx)
            cell.border = cell_border
            if c_idx not in [10, 11, 12, 13, 14]:
                cell.fill = fill_color
        ws_dash.row_dimensions[idx].height = 22

    # Section 2: Obliterated Posts Rehabilitation (Rows 253 onwards)
    sec2_row = 253
    ws_dash.merge_cells(f"A{sec2_row}:Q{sec2_row}")
    sec2_cell = ws_dash.cell(row=sec2_row, column=1, value="SECTION 2: OBLITERATED POST REHABILITATION (NOTIFICATION 1808 - 84 SERVING DISPLACED OFFICERS)")
    sec2_cell.fill = fill_teal
    sec2_cell.font = font_sec_hdr
    sec2_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_dash.row_dimensions[sec2_row].height = 28

    # Repeat headers for Section 2
    sec2_hdr_row = sec2_row + 1
    for c_idx, h in enumerate(dashboard_headers, 1):
        cell = ws_dash.cell(row=sec2_hdr_row, column=c_idx, value=h)
        cell.fill = fill_teal if c_idx not in [10, 11, 12, 13, 14] else (fill_indigo if c_idx in [10, 11, 12] else fill_amber)
        cell.font = font_white_hdr
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
    ws_dash.row_dimensions[sec2_hdr_row].height = 36

    c.execute("""
        SELECT oblit_sl, post_name, establishment, block, district, 
               officer_name, hrms_id, tenure
        FROM obliterated_posts_1808 
        WHERE is_vacant = 'No'
        ORDER BY oblit_sl
    """)
    oblit_serving = c.fetchall()

    for o_idx, r in enumerate(oblit_serving, sec2_hdr_row + 1):
        ws_dash.cell(row=o_idx, column=1, value=r["oblit_sl"]).alignment = Alignment(horizontal="center")
        ws_dash.cell(row=o_idx, column=2, value="Abolished Post Rehab").alignment = Alignment(horizontal="center")
        ws_dash.cell(row=o_idx, column=3, value=r["hrms_id"]).alignment = Alignment(horizontal="center")
        ws_dash.cell(row=o_idx, column=4, value=r["officer_name"]).font = font_bold
        ws_dash.cell(row=o_idx, column=5, value=f"{r['post_name']} ({r['establishment']})")
        ws_dash.cell(row=o_idx, column=6, value=r["block"] or "-")
        ws_dash.cell(row=o_idx, column=7, value=r["district"] or "-")
        ws_dash.cell(row=o_idx, column=8, value=r["tenure"] or "-").alignment = Alignment(horizontal="center")
        ws_dash.cell(row=o_idx, column=9, value="Obliterated Station Rehabilitation").font = font_italic

        ws_dash.cell(row=o_idx, column=10, value="").font = font_bold
        ws_dash.cell(row=o_idx, column=11, value="NO").alignment = Alignment(horizontal="center")
        ws_dash.cell(row=o_idx, column=12, value="").font = font_bold

        status_formula = (
            f'=IF(ISBLANK(J{o_idx}), "PENDING REHABILITATION", '
            f'IF(OR(COUNTIF($J$10:$J$360, J{o_idx})>1, COUNTIF($L$10:$L$360, J{o_idx})>0), '
            f'"⚠️ DUPLICATE ALLOTMENT (" & (COUNTIF($J$10:$J$360, J{o_idx}) + COUNTIF($L$10:$L$360, J{o_idx})) & " OFFICERS SELECTED)", '
            f'IF(K{o_idx}="YES", '
            f'IF(OR(COUNTIF($L$10:$L$360, L{o_idx})>1, COUNTIF($J$10:$J$360, L{o_idx})>0), '
            f'"⚠️ DUPLICATE SU ALLOTMENT (" & (COUNTIF($L$10:$L$360, L{o_idx}) + COUNTIF($J$10:$J$360, L{o_idx})) & " OFFICERS SELECTED)", '
            f'IF(IFERROR(VLOOKUP(L{o_idx}, All_Cadre_Posts_1794!$A$2:$C$2038, 3, FALSE), "")="Occupied", "⚠️ SU COLLISION (DISPLACEMENT TRIGGERED)", "REHABILITATED WITH SU")), '
            f'"REHABILITATED DIRECT")))'
        )
        ws_dash.cell(row=o_idx, column=13, value=status_formula).alignment = Alignment(horizontal="center")

        col_formula = (
            f'=IF(ISBLANK(J{o_idx}), "Awaiting Active Cadre Post", '
            f'IF(COUNTIF($J$10:$J$360, J{o_idx}) + COUNTIF($L$10:$L$360, J{o_idx}) > 1, '
            f'"⚠️ ERROR: Post duplicate allotted across substantive/SU columns! Reassign post.", '
            f'IF(AND(K{o_idx}="YES", NOT(ISBLANK(L{o_idx}))), '
            f'IF(COUNTIF($L$10:$L$360, L{o_idx}) + COUNTIF($J$10:$J$360, L{o_idx}) > 1, '
            f'"⚠️ ERROR: SU post duplicate allotted across substantive/SU columns! Reassign SU post.", '
            f'IF(IFERROR(VLOOKUP(L{o_idx}, All_Cadre_Posts_1794!$A$2:$D$2038, 3, FALSE), "")="Occupied", '
            f'"⚠️ " & IFERROR(VLOOKUP(L{o_idx}, All_Cadre_Posts_1794!$A$2:$D$2038, 4, FALSE), "Incumbent Doctor") & " — MARKED AS ELIGIBLE OFFICER FOR TRANSFER DUE TO DISPLACEMENT", '
            f'"Clear Vacancy / Clean SU")), '
            f'"Direct Cadre Absorption")))'
        )
        ws_dash.cell(row=o_idx, column=14, value=col_formula)

        ws_dash.cell(row=o_idx, column=15, value="Cadre Absorption under Restructuring")
        ws_dash.cell(row=o_idx, column=16, value="Level 16").alignment = Alignment(horizontal="center")
        ws_dash.cell(row=o_idx, column=17, value="Level 16 (Preserved Scale)").alignment = Alignment(horizontal="center")

        fill_color = fill_zebra if o_idx % 2 == 0 else fill_plain
        for c_idx in range(1, 18):
            cell = ws_dash.cell(row=o_idx, column=c_idx)
            cell.border = cell_border
            if c_idx not in [10, 11, 12, 13, 14]:
                cell.fill = fill_color
        ws_dash.row_dimensions[o_idx].height = 22

    # Section 3: Transferable Displaced Officers Pool (Rows 340 onwards)
    sec3_row = sec2_hdr_row + len(oblit_serving) + 2
    ws_dash.merge_cells(f"A{sec3_row}:Q{sec3_row}")
    sec3_cell = ws_dash.cell(row=sec3_row, column=1, value="SECTION 3: TRANSFERABLE DISPLACED QUEUE POOL (OFFICERS DISPLACED DUE TO SERVICE UTILIZATION)")
    sec3_cell.fill = fill_amber
    sec3_cell.font = font_sec_hdr
    sec3_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_dash.row_dimensions[sec3_row].height = 28

    sec3_hdr_row = sec3_row + 1
    for c_idx, h in enumerate(dashboard_headers, 1):
        cell = ws_dash.cell(row=sec3_hdr_row, column=c_idx, value=h)
        cell.fill = fill_amber if c_idx not in [10, 11, 12, 13, 14] else (fill_indigo if c_idx in [10, 11, 12] else fill_navy)
        cell.font = font_white_hdr
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
    ws_dash.row_dimensions[sec3_hdr_row].height = 36

    # 15 Placeholder rows for displaced officers
    for d_idx in range(sec3_hdr_row + 1, sec3_hdr_row + 16):
        d_num = d_idx - sec3_hdr_row
        ws_dash.cell(row=d_idx, column=1, value=d_num).alignment = Alignment(horizontal="center")
        ws_dash.cell(row=d_idx, column=2, value="Transfer due to Displacement").alignment = Alignment(horizontal="center")
        ws_dash.cell(row=d_idx, column=3, value="[Enter HRMS]").alignment = Alignment(horizontal="center")
        ws_dash.cell(row=d_idx, column=4, value=f"Displaced Officer #{d_num} (From SU Collision)").font = font_bold
        ws_dash.cell(row=d_idx, column=5, value="Displaced from Previous Active Post")
        ws_dash.cell(row=d_idx, column=6, value="-")
        ws_dash.cell(row=d_idx, column=7, value="-")
        ws_dash.cell(row=d_idx, column=8, value="-").alignment = Alignment(horizontal="center")
        ws_dash.cell(row=d_idx, column=9, value="Awaiting Rehabilitation Post").font = font_italic

        ws_dash.cell(row=d_idx, column=10, value="").font = font_bold
        ws_dash.cell(row=d_idx, column=11, value="NO").alignment = Alignment(horizontal="center")
        ws_dash.cell(row=d_idx, column=12, value="").font = font_bold

        ws_dash.cell(row=d_idx, column=13, value=f'=IF(ISBLANK(J{d_idx}), "PENDING DISPLACED TRANSFER", IF(OR(COUNTIF($J$10:$J$360, J{d_idx})>1, COUNTIF($L$10:$L$360, J{d_idx})>0), "⚠️ DUPLICATE ALLOTMENT (" & (COUNTIF($J$10:$J$360, J{d_idx}) + COUNTIF($L$10:$L$360, J{d_idx})) & " OFFICERS SELECTED)", "REALLOCATED"))').alignment = Alignment(horizontal="center")
        ws_dash.cell(row=d_idx, column=14, value=f'=IF(ISBLANK(J{d_idx}), "Officer Displaced by Incoming Doctor", IF(COUNTIF($J$10:$J$360, J{d_idx}) + COUNTIF($L$10:$L$360, J{d_idx}) > 1, "⚠️ ERROR: Post duplicate allotted across substantive/SU columns! Reassign post.", "Reallocated to New Post"))')

        ws_dash.cell(row=d_idx, column=15, value="Transfer due to Displacement")
        ws_dash.cell(row=d_idx, column=16, value="Level 16 / 19").alignment = Alignment(horizontal="center")
        ws_dash.cell(row=d_idx, column=17, value="Level 16 / 19").alignment = Alignment(horizontal="center")

        fill_color = fill_zebra if d_idx % 2 == 0 else fill_plain
        for c_idx in range(1, 18):
            cell = ws_dash.cell(row=d_idx, column=c_idx)
            cell.border = cell_border
            if c_idx not in [10, 11, 12, 13, 14]:
                cell.fill = fill_color
        ws_dash.row_dimensions[d_idx].height = 22

    total_dash_rows = sec3_hdr_row + 15

    # =========================================================================
    # DATA VALIDATIONS ON POSTING_DASHBOARD (Dynamic Filter References)
    # =========================================================================
    # 1. Section 1 Roster Candidates: Point to dynamic filtered DD posts column I
    dv_dd = DataValidation(type="list", formula1="'Available_DD_Posts'!$I$2:$I$243", allow_blank=True)
    ws_dash.add_data_validation(dv_dd)
    dv_dd.add("J10:J251")

    # 2. Section 2 Obliterated Posts: Point to dynamic filtered AD vacancies column I
    dv_ad = DataValidation(type="list", formula1=f"'Available_AD_Vacancies'!$I$2:$I${len(ad_rows)+1}", allow_blank=True)
    ws_dash.add_data_validation(dv_ad)
    dv_ad.add(f"J{sec2_hdr_row + 1}:J{sec2_hdr_row + len(oblit_serving)}")

    # 3. Section 3 Displaced Officers: Dynamic filtered DD posts
    dv_disp = DataValidation(type="list", formula1="'Available_DD_Posts'!$I$2:$I$243", allow_blank=True)
    ws_dash.add_data_validation(dv_disp)
    dv_disp.add(f"J{sec3_hdr_row + 1}:J{total_dash_rows}")

    # 4. Toggle SU: NO, YES
    dv_su_toggle = DataValidation(type="list", formula1='"NO,YES"', allow_blank=True)
    ws_dash.add_data_validation(dv_su_toggle)
    dv_su_toggle.add(f"K10:K{total_dash_rows}")

    # 5. SU Post Picker referencing All_Cadre_Posts_1794 Column J (Dynamic Filtered List)
    dv_su = DataValidation(type="list", formula1=f"'All_Cadre_Posts_1794'!$J$2:$J${total_su_rows}", allow_blank=True)
    ws_dash.add_data_validation(dv_su)
    dv_su.add(f"L10:L{total_dash_rows}")

    # =========================================================================
    # CONDITIONAL FORMATTING ON POSTING_DASHBOARD
    # =========================================================================
    # RULE 1: Prominent RED Warning on Column J (Substantive Post Picker) when duplicated across Col J or Col L
    rule_dup_picker_j = FormulaRule(
        formula=['AND(NOT(ISBLANK(J10)), COUNTIF($J$10:$J$360, J10) + COUNTIF($L$10:$L$360, J10) > 1)'],
        stopIfTrue=True,
        fill=fill_red_warning,
        font=font_red_bold
    )
    ws_dash.conditional_formatting.add(f"J10:J{total_dash_rows}", rule_dup_picker_j)

    # RULE 2: Prominent RED Warning on Column L (SU Post Picker) when duplicated across Col L or Col J
    rule_dup_picker_l = FormulaRule(
        formula=['AND(NOT(ISBLANK(L10)), COUNTIF($L$10:$L$360, L10) + COUNTIF($J$10:$J$360, L10) > 1)'],
        stopIfTrue=True,
        fill=fill_red_warning,
        font=font_red_bold
    )
    ws_dash.conditional_formatting.add(f"L10:L{total_dash_rows}", rule_dup_picker_l)

    # RULE 3: Status Column M Formatting
    # A. Duplicate Warning (Vivid Red)
    rule_dup_status = FormulaRule(
        formula=['ISNUMBER(SEARCH("DUPLICATE", M10))'],
        stopIfTrue=True,
        fill=fill_red_warning,
        font=font_red_bold
    )
    # B. SU Collision Warning (Soft Rose)
    rule_collision = FormulaRule(
        formula=['ISNUMBER(SEARCH("COLLISION", M10))'],
        stopIfTrue=True,
        fill=PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid"),
        font=Font(name="Calibri", size=10, bold=True, color="9F1239")
    )
    # C. Direct Placements (Soft Green)
    rule_green = FormulaRule(
        formula=['OR(M10="ALLOTTED DIRECT", M10="REHABILITATED DIRECT", M10="REALLOCATED")'],
        fill=PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid"),
        font=Font(name="Calibri", size=10, bold=True, color="065F46")
    )
    # D. SU Placements (Soft Blue)
    rule_blue = FormulaRule(
        formula=['OR(M10="ALLOTTED WITH SU", M10="REHABILITATED WITH SU")'],
        fill=PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid"),
        font=Font(name="Calibri", size=10, bold=True, color="1E40AF")
    )
    # E. Pending Decisions (Soft Amber)
    rule_pending = FormulaRule(
        formula=['ISNUMBER(SEARCH("PENDING", M10))'],
        fill=PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
        font=Font(name="Calibri", size=10, italic=True, color="92400E")
    )

    ws_dash.conditional_formatting.add(f"M10:M{total_dash_rows}", rule_dup_status)
    ws_dash.conditional_formatting.add(f"M10:M{total_dash_rows}", rule_collision)
    ws_dash.conditional_formatting.add(f"M10:M{total_dash_rows}", rule_green)
    ws_dash.conditional_formatting.add(f"M10:M{total_dash_rows}", rule_blue)
    ws_dash.conditional_formatting.add(f"M10:M{total_dash_rows}", rule_pending)

    # RULE 4: Column N (Displaced Incumbent / Actionable Alert)
    rule_dup_err = FormulaRule(
        formula=['ISNUMBER(SEARCH("ERROR", N10))'],
        stopIfTrue=True,
        fill=fill_red_warning,
        font=font_red_bold
    )
    rule_disp_alert = FormulaRule(
        formula=['ISNUMBER(SEARCH("MARKED AS ELIGIBLE", N10))'],
        fill=PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid"),
        font=Font(name="Calibri", size=10, bold=True, color="9F1239")
    )
    ws_dash.conditional_formatting.add(f"N10:N{total_dash_rows}", rule_dup_err)
    ws_dash.conditional_formatting.add(f"N10:N{total_dash_rows}", rule_disp_alert)

    # RULE 5: KPI Card Red Alert for Duplicates in Rows 5-6
    rule_kpi_dup = FormulaRule(
        formula=['N5>0'],
        fill=fill_red_warning,
        font=Font(name="Calibri", size=18, bold=True, color="991B1B")
    )
    ws_dash.conditional_formatting.add("N5:O6", rule_kpi_dup)

    # =========================================================================
    # CONDITIONAL FORMATTING ON AVAILABLE_DD_POSTS & AVAILABLE_AD_VACANCIES
    # =========================================================================
    fill_avail = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    font_avail = Font(color="047857", bold=True)
    rule_avail = CellIsRule(operator="equal", formula=['"AVAILABLE"'], fill=fill_avail, font=font_avail)

    fill_allot_post = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    font_allot_post = Font(color="64748B")
    rule_allot_post = CellIsRule(operator="equal", formula=['"ALLOTTED"'], fill=fill_allot_post, font=font_allot_post)

    rule_dup = CellIsRule(operator="equal", formula=['"⚠️ DUPLICATE ALLOTMENT!"'], fill=fill_red_warning, font=font_red_bold)

    ws_dd.conditional_formatting.add(f"F2:F{len(dd_rows) + 1}", rule_avail)
    ws_dd.conditional_formatting.add(f"F2:F{len(dd_rows) + 1}", rule_allot_post)
    ws_dd.conditional_formatting.add(f"F2:F{len(dd_rows) + 1}", rule_dup)

    rule_e_dup = FormulaRule(formula=['E2>1'], fill=fill_red_warning, font=font_red_bold)
    ws_dd.conditional_formatting.add(f"E2:E{len(dd_rows) + 1}", rule_e_dup)

    # Available_AD_Vacancies conditional formatting
    ws_ad.conditional_formatting.add(f"F2:F{len(ad_rows) + 1}", rule_avail)
    ws_ad.conditional_formatting.add(f"F2:F{len(ad_rows) + 1}", rule_allot_post)
    ws_ad.conditional_formatting.add(f"F2:F{len(ad_rows) + 1}", rule_dup)
    ws_ad.conditional_formatting.add(f"E2:E{len(ad_rows) + 1}", rule_e_dup)

    # All_Cadre_Posts_1794 conditional formatting
    rule_su_dup = FormulaRule(formula=['H2>1'], fill=fill_red_warning, font=font_red_bold)
    ws_su.conditional_formatting.add(f"H2:H{total_su_rows}", rule_su_dup)

    # =========================================================================
    # TAB 2: Secretariat_Posting_Order (Exact 6 Columns dynamically mirrored)
    # =========================================================================
    ws_order = wb.create_sheet(title="Secretariat_Posting_Order", index=1)
    ws_order.views.sheetView[0].showGridLines = True

    # Government Title Header
    ws_order.merge_cells("A1:F1")
    ws_order.merge_cells("A2:F2")
    ws_order.merge_cells("A3:F3")
    ws_order.merge_cells("A4:F4")

    ws_order.cell(row=1, column=1, value="GOVERNMENT OF WEST BENGAL").font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    ws_order.cell(row=2, column=1, value="Animal Resources Development Department").font = Font(name="Calibri", size=12, bold=True, color="0F766E")
    ws_order.cell(row=3, column=1, value="AR & AH Branch, Prani Sampad Bhawan, LB - 2, Sector - III, Salt Lake, Kolkata - 700 106").font = font_italic
    ws_order.cell(row=4, column=1, value="OFFICIAL POSTING / TRANSFER ORDER SCHEDULE (WBS ROPA RULES 2019)").font = font_bold

    for r_idx in range(1, 5):
        ws_order.cell(row=r_idx, column=1).alignment = Alignment(horizontal="center", vertical="center")
        ws_order.row_dimensions[r_idx].height = 20
    ws_order.row_dimensions[5].height = 8

    # Exact 6 Columns
    sec_headers = [
        "Sl No",
        "Name of the Officers with present place of posting",
        "Present pay level",
        "Transfer by",
        "Place of posting on promotion / transfer/utilization of service",
        "Pay level upon transfer"
    ]
    for c_idx, h in enumerate(sec_headers, 1):
        cell = ws_order.cell(row=6, column=c_idx, value=h)
        cell.fill = fill_navy
        cell.font = font_white_hdr
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
    ws_order.row_dimensions[6].height = 36

    # Dynamically mirror 242 promotees from Posting_Dashboard
    for idx in range(10, 252):
        target_r = idx - 3 # Row 7 to 248
        r_sl = idx - 9

        ws_order.cell(row=target_r, column=1, value=r_sl).alignment = Alignment(horizontal="center", vertical="center")
        
        # Name + Present Posting formula
        name_posting_formula = (
            f'=IF(ISBLANK(Posting_Dashboard!D{idx}), "", '
            f'Posting_Dashboard!D{idx} & " (" & Posting_Dashboard!C{idx} & ") — " & '
            f'Posting_Dashboard!E{idx} & ", " & Posting_Dashboard!F{idx} & ", " & Posting_Dashboard!G{idx})'
        )
        ws_order.cell(row=target_r, column=2, value=name_posting_formula).alignment = Alignment(horizontal="left", vertical="center")
        ws_order.cell(row=target_r, column=3, value=f"=Posting_Dashboard!P{idx}").alignment = Alignment(horizontal="center", vertical="center")
        ws_order.cell(row=target_r, column=4, value=f"=Posting_Dashboard!O{idx}").alignment = Alignment(horizontal="center", vertical="center")

        # Place of posting on promotion / transfer / SU formula
        new_posting_formula = (
            f'=IF(ISBLANK(Posting_Dashboard!J{idx}), "PENDING ALLOTMENT", '
            f'IF(Posting_Dashboard!K{idx}="YES", '
            f'Posting_Dashboard!J{idx} & " [Service Utilization: " & Posting_Dashboard!L{idx} & "]", '
            f'Posting_Dashboard!J{idx}))'
        )
        ws_order.cell(row=target_r, column=5, value=new_posting_formula).alignment = Alignment(horizontal="left", vertical="center")
        ws_order.cell(row=target_r, column=6, value=f"=Posting_Dashboard!Q{idx}").alignment = Alignment(horizontal="center", vertical="center")

        fill_color = fill_zebra if idx % 2 == 0 else fill_plain
        for c_idx in range(1, 7):
            cell = ws_order.cell(row=target_r, column=c_idx)
            cell.border = cell_border
            cell.fill = fill_color
            cell.font = font_bold if c_idx in [1, 2, 5] else font_normal
        ws_order.row_dimensions[target_r].height = 24

    # =========================================================================
    # TAB 6: Instructions_&_Legend
    # =========================================================================
    ws_guide = wb.create_sheet(title="Instructions_&_Legend")
    ws_guide.views.sheetView[0].showGridLines = True

    guide_lines = [
        ("GOVERNMENT OF WEST BENGAL", 16, True, c_navy),
        ("Animal Resources Development Department", 13, True, c_teal),
        ("Smart Posting Decision Board & AI Cadre Cockpit (Google Sheets Interactive Edition)", 12, False, c_slate),
        ("", 10, False, "000000"),
        ("HOW TO USE THE POSTING DASHBOARD:", 12, True, c_navy),
        ("1. Open tab 'Posting_Dashboard': This is your primary interactive decision cockpit.", 10, False, "000000"),
        ("2. Column J ('Eligible Post Picker: Deputy Director'): Click any candidate's cell to open the drop-down list of available DD posts.", 10, False, "000000"),
        ("3. Dynamic Dropdown Reduction: Selected posts automatically vanish from subsequent dropdown menus in Column J and Column L so administrators cannot accidentally re-select them.", 10, False, "000000"),
        ("4. Instant Red Duplicate Warning: If a post is duplicate-allotted in Column J or Column L, the cell turns BRIGHT RED, Column M displays '⚠️ DUPLICATE ALLOTMENT', and the KPI card flags it.", 10, False, "000000"),
        ("5. Column K ('Enable Service Utilization?'): Toggle 'YES' if the officer is to be posted elsewhere under Service Utilization.", 10, False, "000000"),
        ("6. Column L ('Service Utilization Post Picker'): Choose any station across the entire cadre. Once chosen, it automatically vanishes from subsequent SU and substantive pickers!", 10, False, "000000"),
        ("7. Real-Time SU Collision & Displacement Protection: If the station chosen for SU is already occupied, Column M turns RED ('⚠️ SU COLLISION'), and Column N automatically marks the serving doctor as 'MARKED AS ELIGIBLE OFFICER FOR TRANSFER DUE TO DISPLACEMENT'.", 10, False, "000000"),
        ("8. Displaced Officers Pool (Section 3): Any doctor displaced by an SU collision can immediately be assigned a new substantive and SU post in Section 3.", 10, False, "000000"),
        ("9. Dynamic 6-Column Government Order: Open tab 'Secretariat_Posting_Order'. It auto-generates the official notification order live as you make decisions!", 10, False, "000000"),
        ("", 10, False, "000000"),
        ("COLOR CODE LEGEND:", 12, True, c_navy),
        ("🟢 ALLOTTED DIRECT - Officer posted cleanly into clear substantive cadre vacancy (Soft Green).", 10, False, "065F46"),
        ("🔵 ALLOTTED WITH SU - Officer allotted substantive post + Service Utilization to an unblocked station (Soft Blue).", 10, False, "1E40AF"),
        ("🔴 ⚠️ DUPLICATE ALLOTMENT - Same post selected across substantive or SU columns! Cell turns bright red (Red Alert).", 10, False, "991B1B"),
        ("🔴 ⚠️ SU COLLISION - The chosen SU post is occupied! Incumbent officer displaced and flagged for transfer (Soft Rose).", 10, False, "9F1239"),
        ("🟡 PENDING - Awaiting administrative decision (Soft Amber).", 10, False, "92400E"),
        ("⚪ ALLOTTED / BLOCKED - Post has already been taken by an officer in the roster.", 10, False, "475569"),
    ]

    for r_idx, (text, sz, is_bold, color_hex) in enumerate(guide_lines, 1):
        cell = ws_guide.cell(row=r_idx, column=2, value=text)
        cell.font = Font(name="Calibri", size=sz, bold=is_bold, color=color_hex)

    ws_guide.column_dimensions["A"].width = 4
    ws_guide.column_dimensions["B"].width = 110

    # Auto-fit Column Widths across all sheets
    col_widths = {
        "Posting_Dashboard": {
            "A": 6, "B": 20, "C": 14, "D": 28, "E": 32, "F": 22, "G": 18, "H": 14, "I": 28,
            "J": 48, "K": 15, "L": 48, "M": 32, "N": 55, "O": 30, "P": 16, "Q": 32
        },
        "Secretariat_Posting_Order": {
            "A": 8, "B": 55, "C": 18, "D": 30, "E": 60, "F": 32
        },
        "Available_DD_Posts": {
            "A": 10, "B": 52, "C": 35, "D": 18, "E": 25, "F": 24, "G": 28, "H": 36, "I": 52
        },
        "Available_AD_Vacancies": {
            "A": 10, "B": 52, "C": 35, "D": 18, "E": 25, "F": 24, "G": 28, "H": 36, "I": 52
        },
        "All_Cadre_Posts_1794": {
            "A": 55, "B": 40, "C": 18, "D": 34, "E": 18, "F": 14, "G": 10, "H": 25, "I": 24, "J": 55
        }
    }

    for ws in wb.worksheets:
        title = ws.title
        if title in col_widths:
            for col_letter, width in col_widths[title].items():
                ws.column_dimensions[col_letter].width = width

    wb.save(OUTPUT_FILE)
    conn.close()
    print(f"Successfully generated dynamic Google Sheets workbook: {OUTPUT_FILE}")

if __name__ == "__main__":
    build_interactive_dashboard()
