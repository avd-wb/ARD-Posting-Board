import os
import openpyxl
import pypdf
import csv
import pandas as pd

SOURCE_DIR = "/Users/nirmalyaranjansarkar/Projects/AVD"
DEST_DIR = "/Users/nirmalyaranjansarkar/Projects/AVD_AG"

def inspect_transfer_policy():
    pdf_path = os.path.join(SOURCE_DIR, "Transfer Policy/20090219 19.02.2019 Transfer Policy.pdf")
    print(f"\n--- INSPECTING TRANSFER POLICY: {pdf_path} ---")
    if os.path.exists(pdf_path):
        reader = pypdf.PdfReader(pdf_path)
        print(f"Total Pages: {len(reader.pages)}")
        full_text = ""
        for i, page in enumerate(reader.pages):
            txt = page.extract_text() or ""
            full_text += f"\n--- PAGE {i+1} ---\n" + txt
        with open(os.path.join(DEST_DIR, "transfer_policy_text.txt"), "w", encoding="utf-8") as f:
            f.write(full_text)
        print("Extracted first 1500 characters of Transfer Policy:")
        print(full_text[:1500])
    else:
        print("Transfer policy PDF not found!")

def inspect_roster():
    path = os.path.join(SOURCE_DIR, "_00_Sources/01_Verified_Sources /From AD HQ/Revised 50 Point Roster (Only Name) dt. 07-09-2026.xlsx")
    print(f"\n--- INSPECTING REVISED 50 POINT ROSTER: {path} ---")
    df = pd.read_excel(path, sheet_name=0)
    print("Shape:", df.shape)
    print("Columns:", df.columns.tolist())
    print("Head:")
    print(df.head(10))

def inspect_preferences():
    path = os.path.join(SOURCE_DIR, "_00_Sources/01_Verified_Sources /20260906 Posting Preferences.xlsx")
    print(f"\n--- INSPECTING POSTING PREFERENCES (GOOGLE FORM): {path} ---")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Form responses 3"]
    headers = [cell for cell in next(ws.iter_rows(values_only=True))]
    print(f"Total columns: {len(headers)}")
    with open(os.path.join(DEST_DIR, "preference_form_headers.txt"), "w", encoding="utf-8") as f:
        for idx, h in enumerate(headers):
            f.write(f"{idx+1}: {h}\n")
    print("Sample headers (first 25):")
    for idx, h in enumerate(headers[:25]):
        print(f"  {idx+1}: {h}")
    # Also sample 2 rows
    rows = []
    for r in ws.iter_rows(values_only=True):
        rows.append(r)
        if len(rows) >= 3:
            break
    wb.close()

def inspect_sanctioned_posts():
    path = os.path.join(SOURCE_DIR, "04_AVD_Members/06 Vacancies/20260911_AVD_VF-POSTS_Sanctioned_Posts_Incumbents_and_Likely_Vacancies.xlsx")
    print(f"\n--- INSPECTING SANCTIONED POSTS & INCUMBENTS (20260911): {path} ---")
    df = pd.read_excel(path, sheet_name="POSTS")
    print("Shape:", df.shape)
    print("Columns:", df.columns.tolist())
    print("Head (5 rows, first 10 cols):")
    print(df.iloc[:5, :10])

def inspect_cadre_verify():
    path = os.path.join(SOURCE_DIR, "03_ARD_HR/01 Verification data for establishment and Posts from districts/20260911_AVD_CADRE-VERIFY_1794_Posts_Tab_Export.csv")
    print(f"\n--- INSPECTING CADRE VERIFY 1794 POSTS: {path} ---")
    df = pd.read_csv(path)
    print("Shape:", df.shape)
    print("Columns:", df.columns.tolist())
    print("Head (5 rows, first 8 cols):")
    print(df.iloc[:5, :8])

def inspect_gradation_and_bible():
    grad_path = os.path.join(SOURCE_DIR, "03_ARD_HR/ARD_HR/02. Working Outputs/20260908 Gradation List 01092026/20260908_AVD_DDP_Gradation_List_01092026.xlsx")
    print(f"\n--- INSPECTING GRADATION LIST: {grad_path} ---")
    if os.path.exists(grad_path):
        wb = openpyxl.load_workbook(grad_path, read_only=True, data_only=True)
        print("Sheets in Gradation List:", wb.sheetnames)
        ws = wb[wb.sheetnames[0]]
        headers = [c for c in next(ws.iter_rows(values_only=True))]
        print("Gradation List Headers:", headers)
        wb.close()
    
    bible_path = os.path.join(SOURCE_DIR, "03_ARD_HR/ARD_BIBLE_20082026.xlsx")
    print(f"\n--- INSPECTING ARD BIBLE: {bible_path} ---")
    if os.path.exists(bible_path):
        wb = openpyxl.load_workbook(bible_path, read_only=True, data_only=True)
        print("Sheets in ARD Bible:", wb.sheetnames)
        for sname in wb.sheetnames[:3]:
            ws = wb[sname]
            h = [c for c in next(ws.iter_rows(values_only=True))]
            print(f"  Sheet '{sname}' headers: {h[:10]}")
        wb.close()

def inspect_membership():
    zoho_path = os.path.join(SOURCE_DIR, "_00_Sources/Members/20260909 zoho members Contacts.xlsx")
    print(f"\n--- INSPECTING ZOHO MEMBERS: {zoho_path} ---")
    if os.path.exists(zoho_path):
        df = pd.read_excel(zoho_path, nrows=5)
        print("Columns in Zoho Contacts:", df.columns.tolist())
        print("Head:")
        print(df.head(2))

if __name__ == "__main__":
    inspect_transfer_policy()
    inspect_roster()
    inspect_preferences()
    inspect_sanctioned_posts()
    inspect_cadre_verify()
    inspect_gradation_and_bible()
    inspect_membership()
