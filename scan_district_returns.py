import os
import openpyxl
import xlrd
import pandas as pd
import re
from typing import Dict, Any, List

DAHVS_DIR = "/Users/nirmalyaranjansarkar/Projects/AVD/_00_Sources/01_Verified_Sources /From AD HQ DAHVS"

def clean_str(val: Any) -> str:
    if val is None:
        return ""
    s = str(val).strip()
    if s in ["None", "nan"]:
        return ""
    return s

def clean_name(val: Any) -> str:
    s = clean_str(val)
    s = re.sub(r'\(.*?\)', '', s)
    s = s.replace('Dr.', '').replace('Dr', '').strip()
    return s

def clean_hrms(val: Any) -> str:
    if val is None:
        return ""
    s = str(val).strip()
    if s.endswith(".0"):
        s = s[:-2]
    m = re.search(r'\b(19\d{8}|20\d{8})\b', s)
    if m:
        return m.group(1)
    if s in ["None", "nan", "—", "-"]:
        return ""
    return s

def parse_district_files():
    files = [f for f in os.listdir(DAHVS_DIR) if not f.startswith('~$') and not f.startswith('.')]
    print(f"Scanning {len(files)} district status return files in From AD HQ DAHVS...")
    
    records = []
    
    for fname in sorted(files):
        fpath = os.path.join(DAHVS_DIR, fname)
        ext = os.path.splitext(fname)[1].lower()
        
        # parse xlsx
        if ext == ".xlsx":
            try:
                wb = openpyxl.load_workbook(fpath, data_only=True)
                for sname in wb.sheetnames:
                    ws = wb[sname]
                    rows = list(ws.iter_rows(values_only=True))
                    if not rows or len(rows) < 2:
                        continue
                    
                    # Find header row
                    h_idx = -1
                    for idx, r in enumerate(rows[:10]):
                        txt = " ".join([str(c) for c in r if c is not None]).lower()
                        if "name" in txt and ("post" in txt or "designation" in txt or "hrms" in txt or "dob" in txt or "joining" in txt):
                            h_idx = idx
                            break
                    if h_idx != -1:
                        headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[h_idx])]
                        for r in rows[h_idx+1:]:
                            if any(c is not None for c in r):
                                r_dict = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
                                r_dict["_source_file"] = fname
                                r_dict["_source_sheet"] = sname
                                records.append(r_dict)
                wb.close()
            except Exception as e:
                print(f"Error reading {fname}: {e}")
                
        elif ext == ".xls":
            try:
                wb = xlrd.open_workbook(fpath, on_demand=True)
                for sname in wb.sheet_names():
                    sh = wb.sheet_by_name(sname)
                    if sh.nrows < 2:
                        continue
                    h_idx = -1
                    for idx in range(min(10, sh.nrows)):
                        txt = " ".join([str(sh.cell_value(idx, c)) for c in range(sh.ncols)]).lower()
                        if "name" in txt and ("post" in txt or "designation" in txt or "hrms" in txt or "dob" in txt or "joining" in txt):
                            h_idx = idx
                            break
                    if h_idx != -1:
                        headers = [str(sh.cell_value(h_idx, c)).strip() for c in range(sh.ncols)]
                        for row_idx in range(h_idx+1, sh.nrows):
                            r_vals = [sh.cell_value(row_idx, c) for c in range(sh.ncols)]
                            if any(c not in ["", None] for c in r_vals):
                                r_dict = {headers[i]: r_vals[i] for i in range(len(headers))}
                                r_dict["_source_file"] = fname
                                r_dict["_source_sheet"] = sname
                                records.append(r_dict)
            except Exception as e:
                print(f"Error reading XLS {fname}: {e}")

    print(f"Total parsed rows across all district returns: {len(records)}")
    
    # Analyze columns across district returns
    col_counts = {}
    for r in records:
        for k in r.keys():
            if not k.startswith("_"):
                col_counts[k] = col_counts.get(k, 0) + 1
    
    print("\nTop 20 most frequent column headers across district returns:")
    sorted_cols = sorted(col_counts.items(), key=lambda x: x[1], reverse=True)
    for c, cnt in sorted_cols[:20]:
        print(f"  {c}: {cnt} rows")

if __name__ == "__main__":
    parse_district_files()
