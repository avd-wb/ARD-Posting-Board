#!/usr/bin/env python3
"""
generate_final_promotion_242_workbook.py
Generates the authoritative executive Excel deliverable:
Promotion_242_1st_Dradt_20260913_0803.xlsx

Tabs:
1. Promotion_242_Draft - Verified 242 DD promotees on the 50-point roster (15 columns)
2. Master_Posting_Order_310 - Comprehensive 310-officer order (242 promotees + 53 abolished + 15 accommodated)
3. 4_Column_Posting_Order - Official 4-column Government Notification format
4. DD_Vacancy_Balance - Establishment balance sheet of 244 Sanctioned DD posts
5. Discrepancy_Register - 33-point verified audit register
6. Lateral_Field_Adjustments - 13 consequential/field adjustments (Debi Da Rows 313-325)
7. Sources_and_Method - Provenance, methodology, and data sources
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FILE_CLAUDE = "/Users/nirmalyaranjansarkar/Projects/AVD/10_ARD_DD_Promotion_2026/20260913_AVD_DDP_Verified_Posting_Order_and_Discrepancy_Register.xlsx"
FILE_413_MOD = "/Users/nirmalyaranjansarkar/Projects/AVD/10_ARD_DD_Promotion_2026/4.13 am mod 20260913_0012_WB_ARD_Comprehensive_Posting_and_Transfer_Master_Sheet_0.03MB_mb.xlsx"
OUTPUT_LOCAL_AVD = "/Users/nirmalyaranjansarkar/Projects/AVD/10_ARD_DD_Promotion_2026/Promotion_242_1st_Dradt_20260913_0803.xlsx"
OUTPUT_LOCAL_AG = "/Users/nirmalyaranjansarkar/Projects/AVD_AG/Promotion_242_1st_Dradt_20260913_0803.xlsx"

def clean(v):
    return str(v).strip() if v is not None else ""

def build_workbook():
    print("Loading source workbooks...")
    wb_c = openpyxl.load_workbook(FILE_CLAUDE, data_only=True)
    ws_c_master = wb_c["Master_Posting_Order"]
    ws_c_disc = wb_c["Discrepancy_Register"]
    ws_c_vac = wb_c["DD_Vacancy_Balance"]
    ws_c_src = wb_c["Sources_and_Method"]

    wb_413 = openpyxl.load_workbook(FILE_413_MOD, data_only=True)
    ws_413 = wb_413["11_Column_Master_Posting_Order"]

    # Extract Column N comments and bottom additions (Rows 313-325)
    col_n_comments = {}
    for r in range(2, 312):
        comm = ws_413.cell(r, 14).value
        col_n_comments[r] = clean(comm)

    bottom_officers = []
    for r in range(313, 326):
        name = clean(ws_413.cell(r, 3).value)
        if not name:
            name = clean(ws_413.cell(r, 2).value)
        posting = clean(ws_413.cell(r, 8).value)
        comm = clean(ws_413.cell(r, 14).value)
        if name:
            bottom_officers.append({
                "row": r,
                "name": name,
                "target_posting": posting,
                "comments": comm
            })

    # Prepare new workbook
    wb_out = openpyxl.Workbook()
    # remove default sheet
    wb_out.remove(wb_out.active)

    # Styles
    navy_header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    steel_header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    accent_header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    alert_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    crit_alert_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_regular = Font(name="Calibri", size=10)
    font_small = Font(name="Calibri", size=9, italic=True)

    border_thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    headers_15 = [
        "Sl No.",
        "50-Pt Roster Sl No.",
        "Name of the Officer",
        "Present Designation",
        "Present Establishment",
        "Present Block Name (ABAHC/BAHC/BLDO)",
        "Present District",
        "Present Post Description",
        "Present SU (if any)",
        "Transfer / Promotion Basis",
        "Transferred to Substantive Post (DD Level 19)",
        "Service Utilized at (Field / Administrative Post)",
        "Official Remarks",
        "Comments (Debi Da)",
        "Audit & Sanction Status Flags"
    ]

    # --- TAB 1: Promotion_242_Draft ---
    print("Building Tab 1: Promotion_242_Draft...")
    ws1 = wb_out.create_sheet(title="Promotion_242_Draft")
    ws1.views.sheetView[0].showGridLines = True
    ws1.append(headers_15)

    for col_idx in range(1, len(headers_15) + 1):
        cell = ws1.cell(1, col_idx)
        cell.fill = navy_header_fill
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws1.row_dimensions[1].height = 32

    # Populate rows 1 to 242
    for r in range(2, 244): # 242 promotees
        row_out = []
        # sl no (1 to 242)
        row_out.append(ws_c_master.cell(r, 1).value)
        # sl no of 242 promotees
        row_out.append(ws_c_master.cell(r, 2).value)
        # name
        row_out.append(clean(ws_c_master.cell(r, 3).value))
        # present designation
        row_out.append(clean(ws_c_master.cell(r, 4).value))
        # present establishment
        row_out.append(clean(ws_c_master.cell(r, 5).value))
        # block
        row_out.append(clean(ws_c_master.cell(r, 6).value))
        # district
        row_out.append(clean(ws_c_master.cell(r, 7).value))
        # present post
        row_out.append(clean(ws_c_master.cell(r, 8).value))
        # present SU
        row_out.append(clean(ws_c_master.cell(r, 9).value))
        # basis
        row_out.append(clean(ws_c_master.cell(r, 10).value))
        # substantive post
        row_out.append(clean(ws_c_master.cell(r, 11).value))
        # SU post
        row_out.append(clean(ws_c_master.cell(r, 12).value))

        # Split Claude remarks into Official Remarks vs Audit Flags
        rem_full = clean(ws_c_master.cell(r, 13).value)
        parts = rem_full.split(" | ")
        official_rem = parts[0] if parts else ""
        flags = [p for p in parts[1:] if not p.startswith("Review note")]
        flag_str = " | ".join(flags)

        row_out.append(official_rem)
        # Debi Da comments
        row_out.append(col_n_comments.get(r, ""))
        # Audit flags
        row_out.append(flag_str)

        ws1.append(row_out)
        curr_row = ws1.max_row
        ws1.row_dimensions[curr_row].height = 24
        fill_to_use = zebra_fill if curr_row % 2 == 0 else white_fill

        for c in range(1, len(row_out) + 1):
            cell = ws1.cell(curr_row, c)
            cell.font = font_regular
            cell.border = border_thin
            cell.fill = fill_to_use
            if c in [1, 2]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c in [6, 7, 10]:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif c == 14: # Debi Da comments
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                if cell.value:
                    cell.font = font_bold
            elif c == 15: # Audit flags
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                if "CHECK" in str(cell.value):
                    cell.fill = alert_fill
                    cell.font = font_small
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws1.freeze_panes = "D2"

    # --- TAB 2: Master_Posting_Order_310 ---
    print("Building Tab 2: Master_Posting_Order_310...")
    ws2 = wb_out.create_sheet(title="Master_Posting_Order_310")
    ws2.views.sheetView[0].showGridLines = True
    ws2.append(headers_15)

    for col_idx in range(1, len(headers_15) + 1):
        cell = ws2.cell(1, col_idx)
        cell.fill = steel_header_fill
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws2.row_dimensions[1].height = 32

    # Populate all 310 rows
    for r in range(2, 312):
        row_out = []
        row_out.append(ws_c_master.cell(r, 1).value)
        row_out.append(ws_c_master.cell(r, 2).value)
        row_out.append(clean(ws_c_master.cell(r, 3).value))
        row_out.append(clean(ws_c_master.cell(r, 4).value))
        row_out.append(clean(ws_c_master.cell(r, 5).value))
        row_out.append(clean(ws_c_master.cell(r, 6).value))
        row_out.append(clean(ws_c_master.cell(r, 7).value))
        row_out.append(clean(ws_c_master.cell(r, 8).value))
        row_out.append(clean(ws_c_master.cell(r, 9).value))
        row_out.append(clean(ws_c_master.cell(r, 10).value))
        row_out.append(clean(ws_c_master.cell(r, 11).value))
        row_out.append(clean(ws_c_master.cell(r, 12).value))

        rem_full = clean(ws_c_master.cell(r, 13).value)
        parts = rem_full.split(" | ")
        official_rem = parts[0] if parts else ""
        flags = [p for p in parts[1:] if not p.startswith("Review note")]
        flag_str = " | ".join(flags)

        row_out.append(official_rem)
        row_out.append(col_n_comments.get(r, ""))
        row_out.append(flag_str)

        ws2.append(row_out)
        curr_row = ws2.max_row
        ws2.row_dimensions[curr_row].height = 24
        fill_to_use = zebra_fill if curr_row % 2 == 0 else white_fill

        for c in range(1, len(row_out) + 1):
            cell = ws2.cell(curr_row, c)
            cell.font = font_regular
            cell.border = border_thin
            cell.fill = fill_to_use
            if c in [1, 2]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c == 14:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                if cell.value:
                    cell.font = font_bold
            elif c == 15:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                if "CHECK" in str(cell.value):
                    cell.fill = alert_fill
                    cell.font = font_small
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws2.freeze_panes = "D2"

    # --- TAB 3: 4_Column_Posting_Order ---
    print("Building Tab 3: 4_Column_Posting_Order...")
    ws3 = wb_out.create_sheet(title="4_Column_Posting_Order")
    ws3.views.sheetView[0].showGridLines = True

    # Header preamble
    preamble = [
        ["GOVERNMENT OF WEST BENGAL", "", "", ""],
        ["Animal Resources Development Department", "", "", ""],
        ["AR & AH Branch, Prani Sampad Bhawan, LB-2, Sector-III, Salt Lake, Kolkata - 700 106", "", "", ""],
        ["NOTIFICATION (MEMO NO. 1890-AR&AH / DATED 13.09.2026)", "", "", ""],
        ["Schedule I: Promotion to the post of Deputy Director, ARD (Pay Level 19) & Consequential Postings", "", "", ""],
        ["Sl No.", "Name of the Officer with Present Posting", "Place of Posting on Promotion / Transfer (Substantive Post)", "Service Utilized Post (if any) / Remarks"]
    ]
    for row in preamble:
        ws3.append(row)

    # Style preamble
    for r in range(1, 6):
        ws3.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        cell = ws3.cell(r, 1)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.font = Font(name="Calibri", size=12 if r in [1, 4] else 11, bold=True)
        ws3.row_dimensions[r].height = 24

    for c in range(1, 5):
        cell = ws3.cell(6, c)
        cell.fill = accent_header_fill
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws3.row_dimensions[6].height = 28

    # Populate 310 rows
    for r in range(2, 312):
        sl = ws_c_master.cell(r, 1).value
        name = clean(ws_c_master.cell(r, 3).value)
        pres_post = clean(ws_c_master.cell(r, 8).value)
        sub_post = clean(ws_c_master.cell(r, 11).value)
        su_post = clean(ws_c_master.cell(r, 12).value)
        official_rem = clean(ws_c_master.cell(r, 13).value).split(" | ")[0]

        col2 = f"{name}, {pres_post}"
        col4 = su_post if su_post and su_post != "Nil" else official_rem

        ws3.append([sl, col2, sub_post, col4])
        curr_row = ws3.max_row
        ws3.row_dimensions[curr_row].height = 26
        fill_to_use = zebra_fill if curr_row % 2 == 0 else white_fill

        for c in range(1, 5):
            cell = ws3.cell(curr_row, c)
            cell.font = font_regular
            cell.border = border_thin
            cell.fill = fill_to_use
            if c == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws3.freeze_panes = "A7"

    # --- TAB 4: DD_Vacancy_Balance ---
    print("Building Tab 4: DD_Vacancy_Balance...")
    ws4 = wb_out.create_sheet(title="DD_Vacancy_Balance")
    ws4.views.sheetView[0].showGridLines = True

    for r in range(1, ws_c_vac.max_row + 1):
        row_vals = [ws_c_vac.cell(r, c).value for c in range(1, ws_c_vac.max_column + 1)]
        ws4.append(row_vals)
        curr_row = ws4.max_row
        ws4.row_dimensions[curr_row].height = 22

        if r == 1:
            for c in range(1, len(row_vals) + 1):
                cell = ws4.cell(curr_row, c)
                cell.fill = navy_header_fill
                cell.font = font_header
                cell.alignment = Alignment(horizontal="center", vertical="center")
            ws4.row_dimensions[curr_row].height = 28
        else:
            status = clean(ws4.cell(curr_row, 5).value)
            for c in range(1, len(row_vals) + 1):
                cell = ws4.cell(curr_row, c)
                cell.font = font_regular
                cell.border = border_thin
                if status == "OVER-ALLOTTED":
                    cell.fill = crit_alert_fill
                    if c == 5:
                        cell.font = font_bold
                elif status == "Under-allotted":
                    cell.fill = alert_fill
                else:
                    cell.fill = white_fill
                if c in [2, 3, 4]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

    ws4.freeze_panes = "A2"

    # --- TAB 5: Discrepancy_Register ---
    print("Building Tab 5: Discrepancy_Register...")
    ws5 = wb_out.create_sheet(title="Discrepancy_Register")
    ws5.views.sheetView[0].showGridLines = True

    for r in range(1, ws_c_disc.max_row + 1):
        row_vals = [ws_c_disc.cell(r, c).value for c in range(1, ws_c_disc.max_column + 1)]
        ws5.append(row_vals)
        curr_row = ws5.max_row
        ws5.row_dimensions[curr_row].height = 26

        if r == 1:
            for c in range(1, len(row_vals) + 1):
                cell = ws5.cell(curr_row, c)
                cell.fill = navy_header_fill
                cell.font = font_header
                cell.alignment = Alignment(horizontal="center", vertical="center")
            ws5.row_dimensions[curr_row].height = 28
        else:
            sev = clean(ws5.cell(curr_row, 4).value)
            for c in range(1, len(row_vals) + 1):
                cell = ws5.cell(curr_row, c)
                cell.font = font_regular
                cell.border = border_thin
                if sev == "CRITICAL":
                    cell.fill = crit_alert_fill
                elif sev == "HIGH":
                    cell.fill = alert_fill
                else:
                    cell.fill = white_fill
                if c in [1, 4]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    if c == 4:
                        cell.font = font_bold
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws5.freeze_panes = "A2"

    # --- TAB 6: Lateral_Field_Adjustments ---
    print("Building Tab 6: Lateral_Field_Adjustments...")
    ws6 = wb_out.create_sheet(title="Lateral_Field_Adjustments")
    ws6.views.sheetView[0].showGridLines = True

    lat_headers = [
        "Sl No.",
        "Original Master Row",
        "Officer Name",
        "Proposed Target Posting / Direction (Debi Da)",
        "Review Note / Special Comments",
        "Administrative Action Required"
    ]
    ws6.append(lat_headers)
    for c in range(1, len(lat_headers) + 1):
        cell = ws6.cell(1, c)
        cell.fill = accent_header_fill
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws6.row_dimensions[1].height = 28

    for idx, off in enumerate(bottom_officers, start=1):
        action = "Issue executive transfer order in cadre" if off["target_posting"] != "?" else "Clarification required from Debi Da"
        row_vals = [
            idx,
            off["row"],
            off["name"],
            off["target_posting"],
            off["comments"],
            action
        ]
        ws6.append(row_vals)
        curr_row = ws6.max_row
        ws6.row_dimensions[curr_row].height = 24
        fill_to_use = zebra_fill if curr_row % 2 == 0 else white_fill
        for c in range(1, len(row_vals) + 1):
            cell = ws6.cell(curr_row, c)
            cell.font = font_regular
            cell.border = border_thin
            cell.fill = fill_to_use
            if c in [1, 2]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws6.freeze_panes = "A2"

    # --- TAB 7: Sources_and_Method ---
    print("Building Tab 7: Sources_and_Method...")
    ws7 = wb_out.create_sheet(title="Sources_and_Method")
    ws7.views.sheetView[0].showGridLines = True

    for r in range(1, ws_c_src.max_row + 1):
        row_vals = [ws_c_src.cell(r, c).value for c in range(1, ws_c_src.max_column + 1)]
        ws7.append(row_vals)
        curr_row = ws7.max_row
        ws7.row_dimensions[curr_row].height = 22
        for c in range(1, len(row_vals) + 1):
            cell = ws7.cell(curr_row, c)
            cell.font = font_regular
            cell.border = border_thin
            if r == 1:
                cell.fill = navy_header_fill
                cell.font = font_header
            elif cell.value and any(k in str(cell.value) for k in ["WHAT WAS", "OPEN DECISION", "SUMMARY"]):
                cell.font = font_bold
                cell.fill = alert_fill

    # Set optimal column widths across all sheets
    print("Formatting column widths...")
    for ws in wb_out.worksheets:
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            # determine width
            max_len = 0
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len and not cell.coordinate in ws.merged_cells:
                    max_len = len(val)
            optimal_w = min(max(max_len + 3, 12), 48)
            ws.column_dimensions[col_letter].width = optimal_w

    # Custom column adjustments for Tab 1 & 2
    for ws in [ws1, ws2]:
        ws.column_dimensions["A"].width = 8   # Sl No
        ws.column_dimensions["B"].width = 10  # Roster Sl
        ws.column_dimensions["C"].width = 28  # Name
        ws.column_dimensions["D"].width = 30  # Designation
        ws.column_dimensions["E"].width = 32  # Establishment
        ws.column_dimensions["F"].width = 18  # Block
        ws.column_dimensions["G"].width = 20  # District
        ws.column_dimensions["H"].width = 36  # Present Post
        ws.column_dimensions["I"].width = 24  # Present SU
        ws.column_dimensions["J"].width = 16  # Basis
        ws.column_dimensions["K"].width = 38  # Substantive Post
        ws.column_dimensions["L"].width = 38  # SU Post
        ws.column_dimensions["M"].width = 32  # Official Remarks
        ws.column_dimensions["N"].width = 30  # Debi Da Comments
        ws.column_dimensions["O"].width = 45  # Audit & Flags

    # Tab 3 adjustments
    ws3.column_dimensions["A"].width = 8
    ws3.column_dimensions["B"].width = 45
    ws3.column_dimensions["C"].width = 40
    ws3.column_dimensions["D"].width = 40

    # Save to both destinations
    print(f"Saving workbook to {OUTPUT_LOCAL_AVD}...")
    wb_out.save(OUTPUT_LOCAL_AVD)
    print(f"Saving workbook to {OUTPUT_LOCAL_AG}...")
    wb_out.save(OUTPUT_LOCAL_AG)
    print("Workbook creation complete!")

if __name__ == "__main__":
    build_workbook()
