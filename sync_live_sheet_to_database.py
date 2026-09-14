#!/usr/bin/env python3
"""
sync_live_sheet_to_database.py
Synchronizes 337 officers from the latest live Google Sheet into:
1. master_final_order_schedule (337 rows)
2. available_dd_posts (242 allotted, 2 available)
3. cadre_1794_posts (enriched with baseline and post-move vacancy status)
"""

import os
import re
import csv
import sqlite3
import openpyxl
import collections

WORKSPACE = "/Users/nirmalyaranjansarkar/Projects/AVD_AG"
DB_PATH = os.path.join(WORKSPACE, "ard_master_truth.db")
SHEET_PATH = os.path.join(WORKSPACE, "scratch/20260914_AVD_GSH_Review_Sheet_LIVE.xlsx")
LST1_PATH = "/Users/nirmalyaranjansarkar/Projects/ARD PROMOTION/04_LISTS_FROM_SOT/20260914_AVD_LST_1_POSTS_PRESENT_WITH_OCCUPANTS.csv"

def is_staying(o):
    su = (o["su_final"] or "").lower()
    pres = (o["pres_full"] or "").lower()
    blk = (o["block"] or "").lower()
    dist = (o["dist"] or "").lower()
    just = (o["justification"] or "").lower()
    
    if (o["su_final"] or "") in ("", "—", "-"):
        return False
    if any(k in just for k in ["no movement", "keeps working where posted", "place of work (su) is fixed", "seat and the work do not change"]):
        return True
    if "present" in su or "stay" in su:
        return True
    if blk and blk in su and dist in su:
        return True
    return False

def sync_all():
    print("--- 1. Loading Live Google Sheet ---")
    wb = openpyxl.load_workbook(SHEET_PATH, data_only=True)
    ws = wb["REVIEW"]

    officers = []
    for r in range(2, ws.max_row + 1):
        officers.append({
            "sl": int(ws.cell(r, 1).value or (r - 1)),
            "roster_sl": str(ws.cell(r, 2).value or "").strip(),
            "name": str(ws.cell(r, 3).value or "").strip(),
            "hrms": str(ws.cell(r, 4).value or "").strip(),
            "category": str(ws.cell(r, 5).value or "").strip(),
            "desig": str(ws.cell(r, 6).value or "").strip(),
            "estab": str(ws.cell(r, 7).value or "").strip(),
            "block": str(ws.cell(r, 8).value or "").strip(),
            "dist": str(ws.cell(r, 9).value or "").strip(),
            "dor": str(ws.cell(r, 10).value or "").strip(),
            "tenure_left": str(ws.cell(r, 11).value or "").strip(),
            "doj": str(ws.cell(r, 12).value or "").strip(),
            "tenure_in_post": str(ws.cell(r, 13).value or "").strip(),
            "pres_full": str(ws.cell(r, 14).value or "").strip(),
            "pres_su": str(ws.cell(r, 15).value or "").strip(),
            "basis": str(ws.cell(r, 16).value or "").strip(),
            "sub": str(ws.cell(r, 17).value or "").strip(),
            "su": str(ws.cell(r, 18).value or "").strip(),
            "su_final": str(ws.cell(r, 19).value or "").strip(),
            "avd_member": str(ws.cell(r, 20).value or "").strip(),
            "flags": str(ws.cell(r, 21).value or "").strip(),
            "question": str(ws.cell(r, 22).value or "").strip(),
            "options": str(ws.cell(r, 23).value or "").strip(),
            "recommendation": str(ws.cell(r, 24).value or "").strip(),
            "comments": str(ws.cell(r, 25).value or "").strip(),
            "justification": str(ws.cell(r, 26).value or "").strip(),
            "source_check": str(ws.cell(r, 27).value or "").strip()
        })
    print(f"Loaded {len(officers)} officers from LIVE sheet.")

    for o in officers:
        o["is_stay"] = 1 if is_staying(o) else 0

    print("--- 2. Updating master_final_order_schedule ---")
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()

    # Recreate table with comprehensive schema
    cur.execute("DROP TABLE IF EXISTS master_final_order_schedule")
    cur.execute("""
    CREATE TABLE master_final_order_schedule (
        sl_no INTEGER PRIMARY KEY,
        roster_sl TEXT,
        officer_name TEXT,
        clean_name TEXT,
        hrms_id TEXT,
        category TEXT,
        present_designation TEXT,
        present_establishment TEXT,
        present_block TEXT,
        present_district TEXT,
        dor TEXT,
        tenure_left TEXT,
        doj_present_post TEXT,
        tenure_in_post TEXT,
        present_post_full TEXT,
        present_su TEXT,
        transfer_basis TEXT,
        transferred_substantive_post TEXT,
        service_utilized_at TEXT,
        service_utilized_at_final TEXT,
        avd_member TEXT,
        flags TEXT,
        question TEXT,
        options TEXT,
        recommendation TEXT,
        administrative_remarks TEXT,
        comments_directive TEXT,
        justification TEXT,
        source_check TEXT,
        is_stay INTEGER DEFAULT 0
    )
    """)

    order_rows = []
    for o in officers:
        clean = re.sub(r"\s+", " ", re.sub(r"^(dr\.|dr)\s*", "", o["name"], flags=re.I)).strip().lower()
        order_rows.append((
            o["sl"],
            o["roster_sl"],
            o["name"],
            clean,
            o["hrms"],
            o["category"],
            o["desig"],
            o["estab"],
            o["block"],
            o["dist"],
            o["dor"],
            o["tenure_left"],
            o["doj"],
            o["tenure_in_post"],
            o["pres_full"],
            o["pres_su"],
            o["basis"],
            o["sub"],
            o["su"],
            o["su_final"],
            o["avd_member"],
            o["flags"],
            o["question"],
            o["options"],
            o["recommendation"],
            o["comments"] if o["comments"] else ("Stay on SU" if o["is_stay"] else "Transfer/Promotion"),
            o["comments"],
            o["justification"],
            o["source_check"],
            o["is_stay"]
        ))

    cur.executemany("""
    INSERT INTO master_final_order_schedule VALUES (
        ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?
    )
    """, order_rows)
    print(f"Inserted {len(order_rows)} rows into master_final_order_schedule.")

    print("--- 3. Updating available_dd_posts ---")
    cur.execute("UPDATE available_dd_posts SET allotment_status = 'Available', allotted_hrms = NULL, allotted_name = NULL")
    
    promo_officers = [o for o in officers if o["basis"] == "Promotion"]
    allotted_count = 0
    for o in promo_officers:
        h = o["hrms"]
        nm = o["name"]
        target_sub = o["sub"]
        sub_clean = target_sub.lower()
        
        cur.execute("""
            SELECT dd_sl, district, office, post_name 
            FROM available_dd_posts 
            WHERE allotment_status = 'Available'
        """)
        cand_posts = cur.fetchall()
        
        matched_dd_sl = None
        for dd_sl, dist, off, p_name in cand_posts:
            if off and off.lower() in sub_clean:
                matched_dd_sl = dd_sl
                break
            if dist and dist.lower() in sub_clean:
                matched_dd_sl = dd_sl
                break
        
        if not matched_dd_sl and cand_posts:
            matched_dd_sl = cand_posts[0][0]
            
        if matched_dd_sl:
            cur.execute("""
                UPDATE available_dd_posts 
                SET allotment_status = 'Allotted', allotted_hrms = ?, allotted_name = ?
                WHERE dd_sl = ?
            """, (h, nm, matched_dd_sl))
            allotted_count += 1

    cur.execute("SELECT allotment_status, count(*) FROM available_dd_posts GROUP BY allotment_status")
    print("DD post statuses:", cur.fetchall())

    print("--- 4. Enriching cadre_1794_posts with Vacancy Metadata ---")
    cols_meta = [c[1] for c in cur.execute("PRAGMA table_info(cadre_1794_posts)").fetchall()]
    needed_cols = [
        ("post_move_vacancy_status", "TEXT DEFAULT 'NONE'"),
        ("vacated_by_hrms", "TEXT"),
        ("vacated_by_name", "TEXT"),
        ("vacated_by_basis", "TEXT"),
        ("movement_details", "TEXT"),
        ("is_actionable_vacancy", "INTEGER DEFAULT 0"),
        ("is_baseline_vacant", "INTEGER DEFAULT 0"),
        ("is_newly_vacated", "INTEGER DEFAULT 0"),
        ("is_su_retained", "INTEGER DEFAULT 0")
    ]
    for col_name, col_type in needed_cols:
        if col_name not in cols_meta:
            cur.execute(f"ALTER TABLE cadre_1794_posts ADD COLUMN {col_name} {col_type}")

    # Reset metadata
    cur.execute("""
    UPDATE cadre_1794_posts 
    SET post_move_vacancy_status = 'NONE',
        vacated_by_hrms = NULL,
        vacated_by_name = NULL,
        vacated_by_basis = NULL,
        movement_details = NULL,
        is_actionable_vacancy = 0,
        is_baseline_vacant = 0,
        is_newly_vacated = 0,
        is_su_retained = 0
    """)

    # Baseline clear vacancies
    cur.execute("""
    UPDATE cadre_1794_posts
    SET post_move_vacancy_status = 'BASELINE_CLEAR_VACANCY',
        is_baseline_vacant = 1,
        is_actionable_vacancy = 1
    WHERE UPPER(occupancy_status) = 'VACANT'
    """)

    # Load LST_1 for mapping
    lst1_by_pid = {}
    if os.path.exists(LST1_PATH):
        with open(LST1_PATH, encoding="utf-8") as f:
            for r in csv.DictReader(f):
                lst1_by_pid[r["POST_ID"]] = r

    # Process all 1,794 posts
    cur.execute("SELECT id, incumbent_hrms FROM cadre_1794_posts")
    posts = cur.fetchall()

    officers_by_hrms = {o["hrms"]: o for o in officers if o["hrms"]}

    for post_id, inc_hrms in posts:
        pid = f"P{post_id:04d}"
        h = (inc_hrms or "").strip()
        o = officers_by_hrms.get(h)
        if not o and pid in lst1_by_pid:
            l_hrms = (lst1_by_pid[pid].get("HRMS_ID") or "").strip()
            o = officers_by_hrms.get(l_hrms)
            if o:
                h = l_hrms

        if o:
            if o["is_stay"]:
                cur.execute("""
                UPDATE cadre_1794_posts
                SET post_move_vacancy_status = 'SUBSTANTIVE_VACANCY_HELD_ON_SU',
                    is_su_retained = 1,
                    movement_details = ?
                WHERE id = ?
                """, (f"{o['name']} promoted on paper; stays here on SU", post_id))
            else:
                dest = o['sub']
                if o['su_final'] and o['su_final'] not in ("—", "-"):
                    dest += f" (SU at {o['su_final']})"
                cur.execute("""
                UPDATE cadre_1794_posts
                SET post_move_vacancy_status = 'NEWLY_VACATED_BY_MOVE',
                    is_newly_vacated = 1,
                    is_actionable_vacancy = 1,
                    vacated_by_hrms = ?,
                    vacated_by_name = ?,
                    vacated_by_basis = ?,
                    movement_details = ?
                WHERE id = ?
                """, (h, o["name"], o["basis"], f"Vacated by {o['name']} -> {dest}", post_id))

    conn.commit()

    print("\n--- Summary of cadre_1794_posts Post-Move Vacancy Metadata ---")
    counts = cur.execute("""
    SELECT post_move_vacancy_status, count(*) 
    FROM cadre_1794_posts 
    GROUP BY post_move_vacancy_status
    """).fetchall()
    for st, cnt in counts:
        print(f"  {st}: {cnt}")

    actionable_count = cur.execute("SELECT count(*) FROM cadre_1794_posts WHERE is_actionable_vacancy = 1").fetchone()[0]
    print(f"Total Actionable Vacancies (Baseline + Newly Vacated): {actionable_count}")
    
    conn.close()
    print("Database sync completed successfully!")

if __name__ == "__main__":
    sync_all()
