#!/usr/bin/env python3
"""
ingest_dynamic_gradation_list.py

Ingests the latest official Gradation List (brought forward from published list No. 3768-AR&AH)
into ard_master_truth.db table `official_gradation_list`.
Includes:
- Historical 2025 serial (sl_2025)
- Active 2026 consecutive serial (sl_2026), which is NULL for retired / deceased officers
- Status (Serving vs Retired on superannuation vs Deceased)
- is_retired flag (1 for retired/deceased, 0 for serving)
- Officer Name, Clean Name, HRMS ID, Gender, Qualifications, DOB, Age, DOJ, DOR, Category
- Present Posting & Recommended Post
- Cryptographic SHA-256 seal for sacrosanct auditability
"""

import openpyxl
import sqlite3
import hashlib
import re

DB_PATH = 'ard_master_truth.db'
EXCEL_PATH = '/Users/nirmalyaranjansarkar/Projects/AVD/_AI_Generated/03_ARD_HR/ARD_HR/02. Working Outputs/20260908 Gradation List 01092026/20260908_AVD_DDP_Gradation_List_01092026.xlsx'

def clean_str(val):
    if val is None:
        return None
    s = str(val).strip()
    return s if s else None

def parse_int(val):
    if val is None:
        return None
    try:
        return int(val)
    except (ValueError, TypeError):
        return None

def main():
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()

    # Drop and recreate official_gradation_list
    c.execute('DROP TABLE IF EXISTS official_gradation_list')
    c.execute('''
        CREATE TABLE official_gradation_list (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            grade_section TEXT NOT NULL,
            sl_2025 INTEGER,
            sl_2026 INTEGER,
            status_2026 TEXT NOT NULL,
            is_retired INTEGER NOT NULL DEFAULT 0,
            officer_name TEXT NOT NULL,
            clean_name TEXT NOT NULL,
            hrms_id TEXT,
            gender TEXT DEFAULT 'Male',
            qualifications TEXT,
            dob TEXT,
            age TEXT,
            doj TEXT,
            dor TEXT,
            category TEXT,
            avd_member TEXT,
            present_posting TEXT,
            recommended_post TEXT,
            recommendation_reason TEXT,
            remarks TEXT,
            record_sha256 TEXT
        )
    ''')

    # Create index for fast lookups
    c.execute('CREATE INDEX idx_gradation_hrms ON official_gradation_list(hrms_id)')
    c.execute('CREATE INDEX idx_gradation_section ON official_gradation_list(grade_section)')
    c.execute('CREATE INDEX idx_gradation_status ON official_gradation_list(is_retired)')
    c.execute('CREATE INDEX idx_gradation_sl26 ON official_gradation_list(sl_2026)')

    # Load gender map from extended dossier
    gender_map = dict(c.execute('SELECT hrms_id, gender FROM officer_extended_dossier WHERE hrms_id IS NOT NULL').fetchall())

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    sheet = wb['GRADATION_2026']

    inserted = 0
    retired_count = 0
    serving_count = 0

    for row_idx in range(6, sheet.max_row + 1):
        grade_sec = clean_str(sheet.cell(row=row_idx, column=1).value)
        sl_2025 = parse_int(sheet.cell(row=row_idx, column=2).value)
        sl_2026 = parse_int(sheet.cell(row=row_idx, column=3).value)
        status = clean_str(sheet.cell(row=row_idx, column=4).value) or 'Serving'
        officer_name = clean_str(sheet.cell(row=row_idx, column=5).value)
        hrms_id = clean_str(sheet.cell(row=row_idx, column=6).value)

        if not officer_name:
            continue

        qualifications = clean_str(sheet.cell(row=row_idx, column=7).value)
        dob = clean_str(sheet.cell(row=row_idx, column=8).value)
        age = clean_str(sheet.cell(row=row_idx, column=9).value)
        doj = clean_str(sheet.cell(row=row_idx, column=10).value)
        dor = clean_str(sheet.cell(row=row_idx, column=11).value)
        category = clean_str(sheet.cell(row=row_idx, column=12).value) or 'General'
        avd_member = clean_str(sheet.cell(row=row_idx, column=13).value)
        present_posting = clean_str(sheet.cell(row=row_idx, column=16).value)
        rec_post = clean_str(sheet.cell(row=row_idx, column=17).value)
        rec_reason = clean_str(sheet.cell(row=row_idx, column=18).value)
        remarks = clean_str(sheet.cell(row=row_idx, column=22).value)

        clean_name = re.sub(r'^(dr\.?|smt\.?|mr\.?|shri|sri)\s+', '', officer_name.lower()).strip()

        is_ret = 1 if status != 'Serving' else 0
        if is_ret:
            retired_count += 1
            sl_2026 = None  # Ensure active serial is NULL for retired
        else:
            serving_count += 1

        gender = gender_map.get(hrms_id, 'Male')

        # Recalculate SHA-256 hash
        payload = f"{grade_sec}|{sl_2025}|{sl_2026}|{status}|{officer_name}|{hrms_id}|{dob}|{doj}|{dor}|{category}"
        rec_hash = hashlib.sha256(payload.encode('utf-8')).hexdigest()

        c.execute('''
            INSERT INTO official_gradation_list (
                grade_section, sl_2025, sl_2026, status_2026, is_retired,
                officer_name, clean_name, hrms_id, gender, qualifications,
                dob, age, doj, dor, category, avd_member, present_posting,
                recommended_post, recommendation_reason, remarks, record_sha256
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (
            grade_sec, sl_2025, sl_2026, status, is_ret,
            officer_name, clean_name, hrms_id, gender, qualifications,
            dob, age, doj, dor, category, avd_member, present_posting,
            rec_post, rec_reason, remarks, rec_hash
        ))
        inserted += 1

    conn.commit()
    conn.close()

    print(f'Ingestion Complete! Total rows: {inserted}')
    print(f'Serving: {serving_count}, Retired/Left Service: {retired_count}')

if __name__ == '__main__':
    main()
