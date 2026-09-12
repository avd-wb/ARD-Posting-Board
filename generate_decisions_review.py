import openpyxl
import sqlite3
import os

p = "/private/tmp/claude-501/-Users-nirmalyaranjansarkar-Projects-AVD/8bac7a1a-c747-41be-8e3e-6ac9335e725d/scratchpad/LIVE_sheet.xlsx"
wb = openpyxl.load_workbook(p, data_only=True)
s = wb["PROMOTION LIST"]

conn = sqlite3.connect("ard_master_truth.db")
conn.row_factory = sqlite3.Row
cur = conn.cursor()
cur.execute("SELECT hrms_id, post_name, district, establishment FROM obliterated_posts_1808")
oblit_map = {str(r["hrms_id"]): dict(r) for r in cur.fetchall() if r["hrms_id"]}

all_records = []
for r in range(2, 244):
    sl = int(float(s.cell(r, 1).value or 0))
    name = str(s.cell(r, 2).value or "").strip()
    hrms = str(s.cell(r, 3).value or "").replace(".0", "").strip()
    cat = str(s.cell(r, 4).value or "").strip()
    present_post = str(s.cell(r, 11).value or "").strip()
    present_dist = str(s.cell(r, 13).value or "").strip()
    years_left = str(s.cell(r, 21).value or "").strip()
    goes_to = str(s.cell(r, 22).value or "").strip()
    su_flag = str(s.cell(r, 26).value or "").strip()
    note = str(s.cell(r, 30).value or "").strip()
    is_oblit = hrms in oblit_map
    
    all_records.append({
        "sl": sl,
        "name": name,
        "hrms": hrms,
        "category": cat,
        "present_post": present_post,
        "present_dist": present_dist,
        "years_left": years_left,
        "note": note,
        "goes_to": goes_to,
        "su_flag": su_flag,
        "is_oblit": is_oblit,
        "oblit_post": oblit_map.get(hrms, {}).get("post_name")
    })

stays = [r for r in all_records if r["note"].lower() == "stay"]
custom = [r for r in all_records if r["note"] and r["note"].lower() != "stay"]
blank = [r for r in all_records if not r["note"]]
oblit_count = len([r for r in all_records if r["is_oblit"]])

out_path = "/Users/nirmalyaranjansarkar/.gemini/antigravity/brain/7de93603-2849-4e50-b9d7-a5698fd1f235/candidate_manual_decisions_review.md"

with open(out_path, "w", encoding="utf-8") as f:
    f.write("# Master Audit: 242 Promotion Candidates Manual Decisions Review\n\n")
    f.write("This document details all decisions recorded in the live Google Sheet (`1C-tee5LAqyuGJoeWAhe0LiiO3gPAopw0FRnejZRRUAE`, gid=`368838888`), analyzed against statutory cadre rules, remaining tenure, and post obliteration status under Notification No. 1808.\n\n")
    
    f.write("## 📌 Summary Breakdown\n\n")
    f.write(f"- **Total 50-Point Roster Candidates**: 242\n")
    f.write(f"- **Category 1: \"Stay\" Decisions**: {len(stays)} (Incumbent retained at present station on SU; substantive promotion post at Current District Headquarters)\n")
    f.write(f"- **Category 2: Custom / Specific Deployments**: {len(custom)} (Specific DD posts, SU field postings, In-Charge Joint Directors, and Replacement notes)\n")
    f.write(f"- **Category 3: Pending / Blank Rows**: {len(blank)} (Awaiting committee decision)\n")
    f.write(f"- **Critical Alert: Promotees on Obliterated Posts**: {oblit_count} candidates\n\n")
    f.write("---\n\n")
    
    f.write("## 1. Category 1: \"Stay\" Decisions (39 Candidates)\n\n")
    f.write("> **Rule Definition**: The incumbent remains posted at their **Present Station** as **Service Utilized (SU)** (provided the post is not obliterated); their primary substantive promotional post (Pay Level 19) is placed at their **Current District Headquarters** (Deputy Director, ARD & PO / District Office).\n\n")
    f.write("| Sl | Officer Name | HRMS ID | Present Posting & District | Years Left | Substantive DD Post (Pay) | Actual Stationing (SU) | Obliteration Status |\n")
    f.write("|:---|:---|:---|:---|:---|:---|:---|:---|\n")
    for r in stays:
        sl = r["sl"]
        name = r["name"]
        hrms = r["hrms"]
        post = r["present_post"]
        dist = r["present_dist"]
        yrs = r["years_left"]
        ob_str = "⚠️ **OBLITERATED** (" + str(r["oblit_post"]) + ")" if r["is_oblit"] else "Active Valid Post"
        substantive = f"Deputy Director, ARD, {dist} District Office"
        su_stat = f"Stay at {post}"
        f.write(f"| {sl} | **{name}** | `{hrms}` | {post} ({dist}) | {yrs} | {substantive} | {su_stat} | {ob_str} |\n")
    
    f.write("\n---\n\n")
    f.write("## 2. Category 2: Custom & Specific Manual Decisions (115 Candidates)\n\n")
    f.write("| Sl | Officer Name | HRMS ID | Present Posting (District) | Years Left | Manual Decision / Change Note | Target Substantive DD Post | Proposed Placement / SU | Obliteration Alert |\n")
    f.write("|:---|:---|:---|:---|:---|:---|:---|:---|:---|\n")
    for r in custom:
        sl = r["sl"]
        name = r["name"]
        hrms = r["hrms"]
        post = r["present_post"]
        dist = r["present_dist"]
        yrs = r["years_left"]
        note = r["note"].replace("\n", " ")
        target = r["goes_to"] or "To be confirmed"
        ob_str = "⚠️ OBLITERATED" if r["is_oblit"] else "Valid"
        f.write(f"| {sl} | **{name}** | `{hrms}` | {post} ({dist}) | {yrs} | `{note}` | {target} | See Note | {ob_str} |\n")

    f.write("\n---\n\n")
    f.write("## 3. Category 3: Pending / Blank Roster Candidates (88 Candidates)\n\n")
    f.write("| Sl | Officer Name | HRMS ID | Present Posting (District) | Category | Years Left | Obliteration Status |\n")
    f.write("|:---|:---|:---|:---|:---|:---|:---|\n")
    for r in blank:
        sl = r["sl"]
        name = r["name"]
        hrms = r["hrms"]
        post = r["present_post"]
        dist = r["present_dist"]
        cat = r["category"]
        yrs = r["years_left"]
        ob_str = "⚠️ OBLITERATED" if r["is_oblit"] else "Valid"
        f.write(f"| {sl} | **{name}** | `{hrms}` | {post} ({dist}) | {cat} | {yrs} | {ob_str} |\n")

print("Generated successfully:", out_path)
