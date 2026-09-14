import sqlite3
import openpyxl

wb_orig = openpyxl.load_workbook('/Users/nirmalyaranjansarkar/Projects/AVD_AG/comments_final.xlsx', data_only=True)
ws_orig = wb_orig['11_Column_Master_Posting_Order']

conn = sqlite3.connect('ard_master_truth.db')
c = conn.cursor()

c.execute('''
    SELECT sl_no, roster_sl, officer_name, present_designation,
           present_establishment, present_block, present_district, present_post_full,
           present_su, hrms_id
    FROM master_final_order_schedule
    ORDER BY sl_no
''')
corrected_rows = {r[0]: r for r in c.fetchall()}

comparison = []
# Rows 2 to 311 in comments_final
for r_idx in range(2, 312):
    orig_sl = ws_orig.cell(r_idx, 1).value
    orig_rsl = ws_orig.cell(r_idx, 2).value
    orig_name = ws_orig.cell(r_idx, 3).value
    orig_des = ws_orig.cell(r_idx, 4).value or ''
    orig_est = ws_orig.cell(r_idx, 5).value or ''
    orig_blk = ws_orig.cell(r_idx, 6).value or ''
    orig_dist = ws_orig.cell(r_idx, 7).value or ''
    orig_post = ws_orig.cell(r_idx, 8).value or ''
    
    if not orig_sl or orig_sl not in corrected_rows:
        continue
        
    corr = corrected_rows[orig_sl]
    (c_sl, c_rsl, c_name, c_des, c_est, c_blk, c_dist, c_post, c_su, c_hrms) = corr
    
    flaws = []
    if orig_dist != c_dist:
        flaws.append(f"District was '{orig_dist}' -> Corrected to '{c_dist}'")
    if (not orig_blk or orig_blk.strip() == '') and c_blk:
        flaws.append(f"Missing block restored to '{c_blk}'")
    elif orig_blk != c_blk:
        flaws.append(f"Block was '{orig_blk}' -> Corrected to '{c_blk}'")
    if orig_des != c_des:
        flaws.append(f"Designation cleaned from '{orig_des}' -> '{c_des}'")
    if orig_est != c_est:
        flaws.append(f"Establishment clarified from '{orig_est}' -> '{c_est}'")
        
    if flaws:
        comparison.append({
            'sl_no': c_sl,
            'roster_sl': c_rsl,
            'name': c_name,
            'hrms_id': c_hrms,
            'orig_post': orig_post,
            'corr_post': c_post,
            'orig_dist': orig_dist,
            'corr_dist': c_dist,
            'orig_blk': orig_blk,
            'corr_blk': c_blk,
            'flaws': flaws
        })

print(f"Total officers with discrepancies between original comments_final and ground truth: {len(comparison)}")
