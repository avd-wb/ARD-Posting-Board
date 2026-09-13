#!/usr/bin/env python3
"""
generate_final_authoritative_master.py
Builds the definitive, authoritative Promotion and Transfer Master List.
Base Source: debi da final.xlsx (Google Sheet 1ZxTaBpofvb-PvYWIr0Ut3t1MqpOsVbR5 / gid=1976250854).
Strict Enforcement:
- /Users/nirmalyaranjansarkar/Downloads/20260913_1112_AVD_TPV_Draft_Transfer_Order.pdf (17 officers KEPT 100% INTACT)
- Dr. Saravanan E to BAHC Hasnabad, North 24 Parganas
- Dr. Nabadwip Kumar Sarkar to BLDO Basirhat-I, North 24 Parganas
- Dr. Dilip Halder (SC) to Option 4 (SLF Kalyani Substantive + BAHC Kalyani SU)
- Dr. Tarun Kumar Saha Roy (DD Howrah + VO BAHC Shyampur-I SU) & Dr. Manas Kundu (BLDO Kulpi)
- All Debi Da handwritten directives from Column 14 incorporated cleanly.
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SOURCE_EXCEL = "/Users/nirmalyaranjansarkar/Projects/AVD_AG/debi_da_final.xlsx"
OUTPUT_TIMESTAMP = "20260913_1255"
OUTPUT_EXCEL_1 = f"/Users/nirmalyaranjansarkar/Projects/AVD_AG/Promotion_242_Final_List_{OUTPUT_TIMESTAMP}_AG.xlsx"
OUTPUT_EXCEL_2 = f"/Users/nirmalyaranjansarkar/Projects/AVD_AG/Promotion_242_Final_List_{OUTPUT_TIMESTAMP}.AG.xlsx"
OUTPUT_EXCEL_3 = f"/Users/nirmalyaranjansarkar/Projects/AVD/10_ARD_DD_Promotion_2026/Promotion_242_Final_List_{OUTPUT_TIMESTAMP}_AG.xlsx"
OUTPUT_EXCEL_4 = f"/Users/nirmalyaranjansarkar/Projects/AVD/10_ARD_DD_Promotion_2026/Promotion_242_Final_List_{OUTPUT_TIMESTAMP}.AG.xlsx"

# Also maintain standard 1245 alias if requested
ALIAS_1245_1 = "/Users/nirmalyaranjansarkar/Projects/AVD_AG/Promotion_242_Final_List_20260913_1245_AG.xlsx"
ALIAS_1245_2 = "/Users/nirmalyaranjansarkar/Projects/AVD_AG/Promotion_242_Final_List_20260913_1245.AG.xlsx"

def build_master():
    print("Loading source workbook:", SOURCE_EXCEL)
    src_wb = openpyxl.load_workbook(SOURCE_EXCEL, data_only=True)
    src_ws = src_wb['11_Column_Master_Posting_Order']

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws_master = wb.create_sheet(title="Full_Promotion_Transfer_List")

    # Styling definitions
    font_title = Font(name="Calibri", size=13, bold=True, color="1F497D")
    font_sub = Font(name="Calibri", size=11, bold=True, color="333333")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10, bold=False, color="000000")
    font_bold = Font(name="Calibri", size=10, bold=True, color="000000")

    fill_navy = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    fill_slate = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    fill_stay = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fill_trans = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    fill_rehab = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    headers = [
        "Sl No.",
        "Roster Sl (242)",
        "Name of Officer",
        "Present Designation",
        "Present Establishment",
        "Present Block",
        "Present District",
        "Present Post (Full Spec)",
        "Present SU (if any)",
        "Transfer Basis",
        "Transferred Substantive Post",
        "Service Utilized At (SU)",
        "Administrative Remarks",
        "Comments (Debi Da Directive)"
    ]

    for c_idx, h in enumerate(headers, 1):
        cell = ws_master.cell(1, c_idx, h)
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws_master.row_dimensions[1].height = 32

    # 17 Special Officers from 1112 PDF (Strictly Preserved)
    tpv_17 = {
        "prasanta kumar bera": {
            "sub": "Assistant Director, ARD (Veterinary), Directorate Headquarter, Kolkata",
            "su": "Nil",
            "rem": "Transferred as Assistant Director, ARD (Veterinary), Directorate Headquarter, Kolkata",
            "basis": "Displacement due to post abolition"
        },
        "pradip pati": {
            "sub": "Assistant Director, ARD (Veterinary), Directorate Headquarter, Kolkata",
            "su": "Nil",
            "rem": "Transferred as Assistant Director, ARD (Veterinary), Directorate Headquarter, Kolkata",
            "basis": "Displacement due to post abolition"
        },
        "sukanta roy": {
            "sub": "Assistant Director, ARD (Veterinary), Directorate Headquarter, Kolkata",
            "su": "Nil",
            "rem": "Transferred as Assistant Director, ARD (Veterinary), Directorate Headquarter, Kolkata",
            "basis": "Displacement due to post abolition"
        },
        "debi prasad nandi": {
            "sub": "Assistant Director, ARD, District Office, North 24 Parganas",
            "su": "Assistant Director, ARD (Veterinary), Directorate Headquarter, Kolkata",
            "rem": "Transferred as Assistant Director, ARD, District Office, North 24 Parganas; service utilized as Assistant Director, ARD (Veterinary), Directorate Headquarter, Kolkata",
            "basis": "Displacement due to post abolition"
        },
        "nirmalya ranjan sarkar": {
            "sub": "Assistant Director, ARD, District Office, Hooghly",
            "su": "Assistant Director, ARD (Veterinary), Directorate Headquarter, Kolkata",
            "rem": "Post abolished; rehabilitated in the active cadre as Assistant Director, ARD, District Office, Hooghly; service utilized as Assistant Director, ARD (Veterinary), Directorate Headquarter, Kolkata",
            "basis": "Displacement due to post abolition"
        },
        "puspendu panja": {
            "sub": "Assistant Director, ARD (Veterinary), Directorate Headquarter, Kolkata",
            "su": "Nil",
            "rem": "Transferred as Assistant Director, ARD (Veterinary), Directorate Headquarter, Kolkata",
            "basis": "Displacement due to post abolition"
        },
        "soma das": {
            "sub": "Deputy Director, ARD, O/O the DAH & VS, W.B, Directorate Headquarters",
            "su": "Nil",
            "rem": "Promoted to Deputy Director, ARD, Directorate Headquarters",
            "basis": "Promotion"
        },
        "kartick chandra roy": {
            "sub": "Deputy Director, ARD, O/O the DAH & VS, W.B, Directorate Headquarters",
            "su": "Nil",
            "rem": "Promoted to Deputy Director, ARD, Directorate Headquarters",
            "basis": "Promotion"
        },
        "banibrata nayek": {
            "sub": "Assistant Director, ARD (Veterinary), Directorate Headquarter, Kolkata",
            "su": "Nil",
            "rem": "Transferred as Assistant Director, ARD (Veterinary), Directorate Headquarter, Kolkata",
            "basis": "Displacement due to promotee accommodation"
        },
        "samir patra": {
            "sub": "Deputy Director, ARD, O/O the DAH & VS, W.B, Directorate Headquarters",
            "su": "Nil",
            "rem": "Promoted to Deputy Director, ARD, Directorate Headquarters",
            "basis": "Promotion"
        },
        "chayan bhattacharya": {
            "sub": "Deputy Director, ARD, Haringhata Farm, Nadia",
            "su": "Nil",
            "rem": "Promoted to Deputy Director, ARD, Haringhata Farm, Nadia",
            "basis": "Promotion"
        },
        "shuvendu halder": {
            "sub": "Block Livestock Development Officer, Sub-Divisional and Block Level Set up of South 24 Parganas, Mathurapur-I, South 24 Parganas",
            "su": "Nil",
            "rem": "Transferred as Block Livestock Development Officer, Mathurapur-I, South 24 Parganas in lieu of Dr. Kartick Chandra Roy",
            "basis": "Displacement due to promotee accommodation"
        },
        "madhusudan mukherjee": {
            "sub": "Assistant Director, ARD (Veterinary), Directorate Headquarter, Kolkata",
            "su": "Manager (HR), WBLDCL, Headquarters",
            "rem": "Transferred as Assistant Director, ARD (Veterinary), Directorate Headquarter, Kolkata; service utilized as Manager (HR), WBLDCL, Headquarters",
            "basis": "Displacement due to promotee accommodation"
        },
        "partha sarathi chattopadhyay": {
            "sub": "Block Livestock Development Officer, Sub-Divisional and Block Level Set up of Dakshin Dinajpur District, Kushmandi, Dakshin Dinajpur",
            "su": "Nil",
            "rem": "Transferred as Block Livestock Development Officer, Kushmandi, Dakshin Dinajpur",
            "basis": "Displacement due to promotee accommodation"
        },
        "dipak dey": {
            "sub": "Assistant Director, ARD (Veterinary), Directorate Headquarter, Kolkata",
            "su": "Marketing, WBLDCL, Headquarters",
            "rem": "Transferred as Assistant Director, ARD (Veterinary), Directorate Headquarter, Kolkata; service utilized as Marketing, WBLDCL, Headquarters",
            "basis": "Displacement due to promotee accommodation"
        },
        "santanu nandi": {
            "sub": "Block Livestock Development Officer, Sub-Divisional and Block Level Set up of North 24 Parganas District, Bongaon, North 24 Parganas",
            "su": "Nil",
            "rem": "Service Utilization withdrawn and posted as Block Livestock Development Officer, Bongaon, North 24 Parganas",
            "basis": "Displacement due to promotee accommodation"
        },
        "sumit chowdhury": {
            "sub": "Assistant Director, ARD (Veterinary Research & Investigation), I.A.H. & V.B., Belgachia, Kolkata",
            "su": "Nil",
            "rem": "Transferred as Assistant Director, ARD (VR&I), I.A.H. & V.B., Belgachia, Kolkata",
            "basis": "Displacement due to promotee accommodation"
        }
    }

    master_records = []
    seen_names = set()

    # Pass 1: Read all rows from debi_da_final.xlsx (rows 2 to 311)
    for r in range(2, 312):
        sl = src_ws.cell(r, 1).value
        rsl = src_ws.cell(r, 2).value
        name = src_ws.cell(r, 3).value
        pdes = src_ws.cell(r, 4).value
        pest = src_ws.cell(r, 5).value
        pblk = src_ws.cell(r, 6).value
        pdist = src_ws.cell(r, 7).value
        ppost = src_ws.cell(r, 8).value
        psu = src_ws.cell(r, 9).value
        tbasis = src_ws.cell(r, 10).value
        sub = src_ws.cell(r, 11).value
        su = src_ws.cell(r, 12).value
        rem = src_ws.cell(r, 13).value
        deb = src_ws.cell(r, 14).value

        if not name:
            continue

        name_str = str(name).strip()
        deb_str = str(deb).strip() if deb else ""
        sub_str = str(sub).strip() if sub else ""
        su_str = str(su).strip() if su else ""
        rem_str = str(rem).strip() if rem else ""
        ppost_str = str(ppost).strip() if ppost else ""
        pdist_str = str(pdist).strip() if pdist else ""
        pblk_str = str(pblk).strip() if pblk else ""

        try:
            rsl_int = int(float(rsl)) if rsl is not None and str(rsl).strip() not in ['-', ''] else None
        except Exception:
            rsl_int = None

        final_sub = sub_str
        final_su = su_str
        final_rem = rem_str
        final_basis = str(tbasis).strip() if tbasis else "Promotion"

        # Check if matched in TPV 17
        matched_tpv = None
        for k, v in tpv_17.items():
            if k in name_str.lower():
                matched_tpv = v
                break

        if matched_tpv:
            final_sub = matched_tpv['sub']
            final_su = matched_tpv['su']
            final_rem = matched_tpv['rem']
            final_basis = matched_tpv['basis']
        elif "saravan" in name_str.lower():
            # Dr. Saravanan E
            final_sub = "Veterinary Officer, BAHC, Sub-Divisional and Block Level Set up of North 24 Parganas District, Hasnabad, North 24 Parganas"
            final_su = "Nil"
            final_rem = "Transferred as Veterinary Officer, BAHC, Hasnabad, North 24 Parganas"
            final_basis = "Transfer"
        elif "nabadwip kumar sarkar" in name_str.lower():
            # Dr. Nabadwip Kumar Sarkar
            final_sub = "Block Livestock Development Officer, Sub-Divisional and Block Level Set up of North 24 Parganas District, Basirhat-I, North 24 Parganas"
            final_su = "Nil"
            final_rem = "Transferred as Block Livestock Development Officer, Basirhat-I, North 24 Parganas"
            final_basis = "Displacement due to post abolition"
        elif rsl_int == 193 or "tarun kumar saha roy" in name_str.lower():
            final_sub = "Deputy Director, ARD, O/O the JD ARD, Howrah"
            final_su = "Veterinary Officer, BAHC, Sub-Divisional and Block Level Set up of Howrah, Shyampur-I, Howrah"
            final_rem = "Promoted to Deputy Director, ARD, Howrah; service utilized as Veterinary Officer, BAHC, Shyampur-I, Howrah"
            final_basis = "Promotion"
        elif rsl_int == 132 or "dilip halder" in name_str.lower():
            # Option 4 (Kalyani Field Deployment)
            final_sub = "Deputy Director, ARD, State Livestock Farm, Kalyani"
            final_su = "Veterinary Officer, BAHC, Sub-Divisional and Block Level Set up of Nadia District, Kalyani, Nadia"
            final_rem = "Promoted to Deputy Director, ARD, State Livestock Farm, Kalyani; service utilized as Veterinary Officer, BAHC, Kalyani, Nadia"
            final_basis = "Promotion"
        elif rsl_int == 242 or "basudev datta" in name_str.lower():
            final_sub = "Deputy Director, ARD, Quarantine Station, Naxalbari, Siliguri"
            final_su = "Block Livestock Development Officer, Sub-Divisional and Block Level Set up of Coochbehar District, Tufanganj-II, Coochbehar"
            final_rem = "Stay at present station on SU (Pay Level 19 at District HQ)"
            final_basis = "Promotion"
        elif rsl_int == 206 or "masur ali" in name_str.lower():
            final_sub = "Deputy Director, ARD, State Poultry Farm, Domkal, Murshidabad"
            final_su = "Nil"
            final_rem = "Promoted to Deputy Director, ARD, Murshidabad"
            final_basis = "Promotion"
        elif rsl_int == 199 or "tapan kumar sur" in name_str.lower():
            final_sub = "Deputy Director, ARD, District Office, Siliguri"
            final_su = "Nil"
            final_rem = "Promoted to Deputy Director, ARD, Siliguri"
            final_basis = "Promotion"
        elif rsl_int == 210 or "sushil kr. baskey" in name_str.lower():
            final_sub = "Deputy Director, ARD, O/O the JD ARD, Purulia"
            final_su = "Nil"
            final_rem = "Promoted to Deputy Director, ARD, Purulia"
            final_basis = "Promotion"
        elif rsl_int == 241 or "palash hansda" in name_str.lower():
            final_sub = "Deputy Director, ARD, District Office, Paschim Medinipur"
            final_su = "Nil"
            final_rem = "Promoted to Deputy Director, ARD, Paschim Medinipur"
            final_basis = "Promotion"
        elif rsl_int == 236 or "kamal sinha" in name_str.lower():
            final_sub = "Deputy Director, ARD, O/O the JD ARD, Siliguri"
            final_su = "Block Livestock Development Officer, Sub-Divisional and Block Level Set up of Purulia, Barabazar, Purulia"
            final_rem = "Promoted to Deputy Director, ARD, Siliguri; service utilized as BLDO Barabazar, Purulia"
            final_basis = "Promotion"
        elif rsl_int == 237 or "chinmoy mitra" in name_str.lower():
            final_sub = "Deputy Director, ARD, State Poultry Farm, Mohitnagar, Jalpaiguri"
            final_su = "Block Livestock Development Officer, Sub-Divisional and Block Level Set up of Murshidabad District, Domkal, Murshidabad"
            final_rem = "Promoted to Deputy Director, ARD, Jalpaiguri; service utilized as BLDO Domkal, Murshidabad"
            final_basis = "Promotion"
        elif rsl_int == 238 or "sisir roy" in name_str.lower():
            final_sub = "Deputy Director, ARD, PMC, Kalimpong"
            final_su = "Block Livestock Development Officer, Sub-Divisional and Block Level Set up of Nadia District, Nadia"
            final_rem = "Promoted to Deputy Director, ARD, Kalimpong; service utilized as BLDO in Nadia"
            final_basis = "Promotion"
        elif rsl_int is not None and rsl_int <= 242:
            if deb_str.lower() in ['stay', 's']:
                if 'deputy director' in sub_str.lower() or 'ddard' in sub_str.lower():
                    final_sub = sub_str
                else:
                    final_sub = f"Deputy Director, ARD, District Office, {pdist_str}"
                final_su = ppost_str
                final_rem = "Stay at present station on SU (Pay Level 19 at District HQ)"
            elif deb_str.lower().startswith('ddard') or deb_str.lower().startswith('dd') or deb_str.lower().startswith('deputy director'):
                cleaned = deb_str.replace("DDARD,", "Deputy Director, ARD,").replace("DDARD", "Deputy Director, ARD,")
                cleaned = cleaned.replace("DD,", "Deputy Director, ARD,").replace("DD ", "Deputy Director, ARD, ")
                cleaned = cleaned.replace("Dte. HQ", "O/O the DAH & VS, W.B, Directorate Headquarters")
                cleaned = cleaned.replace("IAH&VB, Kolkata", "O/O the Additional Director, ARD, I.A.H. & V.B., (R. & T.)")
                final_sub = cleaned
                final_su = "Nil"
                final_rem = f"Promoted to {cleaned}"
            elif any(deb_str.lower().startswith(k) for k in ['su as', 'bldo', 'vo', 'bahc', 'abahc', 'sahc']):
                cleaned_su = deb_str.replace("SU as ", "").replace("su as ", "")
                cleaned_su = cleaned_su.replace("Burdwan Murshidabad", "Burwan, Murshidabad")
                cleaned_su = cleaned_su.replace("Muhhadbazar Birbhum", "Md. Bazar, Birbhum")
                cleaned_su = cleaned_su.replace("Hirband Bankura", "Hirbandh, Bankura")
                cleaned_su = cleaned_su.replace("Toofangunj", "Tufanganj").replace("toofangunj", "Tufanganj")
                cleaned_su = cleaned_su.replace("Kotchila", "Kotsila").replace("Shyamput_I", "Shyampur-I")
                final_su = cleaned_su
                if 'deputy director' in sub_str.lower() or 'ddard' in sub_str.lower():
                    final_sub = sub_str
                else:
                    final_sub = f"Deputy Director, ARD, District Office, {pdist_str}"
                final_rem = f"Promoted to {final_sub}; service utilized as {final_su}"
            elif deb_str in ['?', '']:
                final_sub = sub_str if sub_str else f"Deputy Director, ARD, District Office, {pdist_str}"
                final_su = su_str if su_str else "Nil"
                final_rem = rem_str if rem_str else f"Promoted to {final_sub}"
            else:
                if deb_str in ['Malda', 'Coochbehar']:
                    final_sub = f"Deputy Director, ARD, District Office, {deb_str}"
                    final_su = "Nil"
                    final_rem = f"Promoted to Deputy Director, ARD, {deb_str}"
                elif 'polyclinic behala' in deb_str.lower():
                    final_sub = "Deputy Director, ARD, Veterinary Polyclinic, Behala, South 24 Parganas"
                    final_su = "Nil"
                    final_rem = "Promoted to Deputy Director, ARD, Veterinary Polyclinic, Behala, South 24 Parganas"
                elif 'north bengal' in deb_str.lower():
                    final_sub = "Deputy Director, ARD, O/O the Additional Director, ARD, Set up of North Bengal"
                    final_su = "Nil"
                    final_rem = "Promoted to Deputy Director, ARD, Set up of North Bengal"
                elif 'nearby ddard' in deb_str.lower():
                    final_sub = "Deputy Director, ARD, District Office, Paschim Medinipur"
                    final_su = "Nil"
                    final_rem = "Promoted to Deputy Director, ARD, Paschim Medinipur"
                elif 'kantapukur' in deb_str.lower():
                    final_sub = "Deputy Director, ARD, State Poultry Farm, Kantapukur, Howrah"
                    final_su = "Nil"
                    final_rem = "Promoted to Deputy Director, ARD, State Poultry Farm, Kantapukur, Howrah"
                elif 'haringhata' in deb_str.lower() and 'management' in deb_str.lower():
                    final_sub = "Deputy Director, ARD, Haringhata Farm, Nadia"
                    final_su = "Assistant Director, ARD, (Management), Haringhata Farm, Nadia"
                    final_rem = "Promoted to Deputy Director, ARD, Haringhata Farm; service utilized as Assistant Director, ARD (Management)"
                elif 'birbhum bldo' in deb_str.lower():
                    final_sub = "Deputy Director, ARD, District Office, Birbhum"
                    final_su = "Block Livestock Development Officer, Sub-Divisional and Block Level Set up of Birbhum"
                    final_rem = "Promoted to Deputy Director, ARD, Birbhum; service utilized as BLDO in Birbhum"
                elif 'howrah bldo' in deb_str.lower():
                    final_sub = "Deputy Director, ARD, District Office, Howrah"
                    final_su = "Block Livestock Development Officer, Sub-Divisional and Block Level Set up of Howrah"
                    final_rem = "Promoted to Deputy Director, ARD, Howrah; service utilized as BLDO in Howrah"
                elif 'et lab' in deb_str.lower():
                    final_sub = "Deputy Director, ARD, Haringhata Farm, Nadia"
                    final_su = "Station Director, Embryo Transfer Laboratory, PBGSBS, Haringhata Farm, Nadia"
                    final_rem = "Promoted to Deputy Director, ARD, Haringhata Farm; service utilized as Station Director, ET Lab, PBGSBS"
                elif 'murshidabad' in deb_str.lower() and 'bldo' in deb_str.lower():
                    final_sub = "Deputy Director, ARD, District Office, Murshidabad"
                    final_su = "Block Livestock Development Officer, Sub-Divisional and Block Level Set up of Murshidabad"
                    final_rem = "Promoted to Deputy Director, ARD, Murshidabad; service utilized as BLDO in Murshidabad"
                elif 'purulia' in deb_str.lower() and 'bankura' in deb_str.lower():
                    final_sub = "Deputy Director, ARD, O/O the JD ARD, Purulia"
                    final_su = "Block Livestock Development Officer, Joypur, Purulia"
                    final_rem = "Promoted to Deputy Director, ARD, Purulia; service utilized as BLDO Joypur, Purulia"
                else:
                    final_sub = sub_str
                    final_su = su_str
                    final_rem = rem_str

        record = {
            'rsl': rsl_int if rsl_int is not None else "-",
            'name': name_str,
            'pdes': str(pdes or "").strip(),
            'pest': str(pest or "").strip(),
            'pblk': pblk_str,
            'pdist': pdist_str,
            'ppost': ppost_str,
            'psu': str(psu or "Nil").strip(),
            'basis': final_basis,
            'sub': final_sub,
            'su': final_su if final_su else "Nil",
            'rem': final_rem,
            'deb': deb_str
        }
        master_records.append(record)
        seen_names.add(name_str.lower())

    # Pass 2: Append subsequent displacement / chain transfer officers
    extra_officers = [
        {
            'rsl': "-",
            'name': "Dr. Falguni Chakraborty",
            'pdes': "Block Livestock Development Officer",
            'pest': "Sub-Divisional and Block Level Set up of Hooghly",
            'pblk': "Goghat-I",
            'pdist': "Hooghly",
            'ppost': "Block Livestock Development Officer, Sub-Divisional and Block Level Set up of Hooghly, Goghat-I, Hooghly",
            'psu': "Nil",
            'basis': "Transfer",
            'sub': "Block Livestock Development Officer, Sub-Divisional and Block Level Set up of Murshidabad District, Nabagram, Murshidabad",
            'su': "Nil",
            'rem': "Transferred as Block Livestock Development Officer, Nabagram, Murshidabad",
            'deb': "Murshidabad"
        },
        {
            'rsl': "-",
            'name': "Dr. Badal Chandra Das",
            'pdes': "Block Livestock Development Officer",
            'pest': "Sub-Divisional and Block Level Set up of Uttar Dinajpur District",
            'pblk': "Goalpokher-II",
            'pdist': "Uttar Dinajpur",
            'ppost': "Block Livestock Development Officer, Sub-Divisional and Block Level Set up of Uttar Dinajpur District, Goalpokher-II, Uttar Dinajpur",
            'psu': "Nil",
            'basis': "Transfer",
            'sub': "Block Livestock Development Officer, Sub-Divisional and Block Level Set up of Murshidabad District, Berhampore, Murshidabad",
            'su': "Nil",
            'rem': "Transferred as Block Livestock Development Officer, Berhampore, Murshidabad",
            'deb': "Murshidabad"
        },
        {
            'rsl': "-",
            'name': "Dr. Rajib Kanti Saha",
            'pdes': "Veterinary Officer",
            'pest': "O/O the Deputy Director, ARD & Parishad Officer",
            'pblk': "",
            'pdist': "Cooch Behar",
            'ppost': "Veterinary Officer, O/O the Deputy Director, ARD & Parishad Officer, Cooch Behar",
            'psu': "Nil",
            'basis': "Transfer",
            'sub': "Veterinary Officer, SAHC, Sub-Divisional and Block Level Set up of South 24 Parganas, Canning-I, South 24 Parganas",
            'su': "Nil",
            'rem': "Transferred as Veterinary Officer, SAHC, Canning-I, South 24 Parganas",
            'deb': "SAHC Canning"
        },
        {
            'rsl': "-",
            'name': "Dr. Sanjay Sheet",
            'pdes': "Veterinary Officer",
            'pest': "Sub-Divisional and Block Level Set up of Dakshin Dinajpur District",
            'pblk': "Tapan",
            'pdist': "Dakshin Dinajpur",
            'ppost': "Veterinary Officer, Sub-Divisional and Block Level Set up of Dakshin Dinajpur District, Tapan, Dakshin Dinajpur",
            'psu': "Nil",
            'basis': "Transfer",
            'sub': "Veterinary Officer, Veterinary Polyclinic, Balurghat, Dakshin Dinajpur",
            'su': "Nil",
            'rem': "Transferred as Veterinary Officer, Veterinary Polyclinic, Balurghat, Dakshin Dinajpur",
            'deb': "as per choice"
        },
        {
            'rsl': "-",
            'name': "Dr. Ritesh Biswas",
            'pdes': "Assistant Director, ARD",
            'pest': "I.A.H. & V.B., (R. & T.), Belgachia",
            'pblk': "",
            'pdist': "Kolkata",
            'ppost': "Assistant Director, ARD, I.A.H. & V.B., (R. & T.), Belgachia, Kolkata",
            'psu': "Nil",
            'basis': "Transfer",
            'sub': "Assistant Director, ARD, Vety. Polyclinic, Bardhaman, Purba Bardhaman",
            'su': "Nil",
            'rem': "Transferred as Assistant Director, ARD, Vety. Polyclinic, Bardhaman, Purba Bardhaman",
            'deb': "burdwan polyclinic"
        },
        {
            'rsl': "-",
            'name': "Dr. Kuntal Roy",
            'pdes': "Veterinary Officer",
            'pest': "O/O the Block Livestock Development Officer",
            'pblk': "Hilly",
            'pdist': "Dakshin Dinajpur",
            'ppost': "Veterinary Officer, O/O the Block Livestock Development Officer, Hilly, Dakshin Dinajpur",
            'psu': "Nil",
            'basis': "Transfer",
            'sub': "Block Livestock Development Officer, Sub-Divisional and Block Level Set up of Coochbehar District, Dinhata-I, Coochbehar",
            'su': "Nil",
            'rem': "Transferred as Block Livestock Development Officer, Dinhata-I, Coochbehar",
            'deb': "Any post in coochbehar district"
        },
        {
            'rsl': "-",
            'name': "Dr. Manas Kundu",
            'pdes': "Veterinary Officer",
            'pest': "Office of the Block Livestock Development Officer, Shyampur-I",
            'pblk': "Shyampur-I",
            'pdist': "Howrah",
            'ppost': "Veterinary Officer, Office of the Block Livestock Development Officer, Shyampur-I, Shyampur-I, Howrah",
            'psu': "Nil",
            'basis': "Transfer",
            'sub': "Block Livestock Development Officer, Sub-Divisional and Block Level Set up of South 24 Parganas, Kulpi, South 24 Parganas",
            'su': "Nil",
            'rem': "Transferred as Block Livestock Development Officer, Sub-Divisional and Block Level Set up of South 24 Parganas, Kulpi, South 24 Parganas",
            'deb': "BLDO Kulpi"
        },
        {
            'rsl': "-",
            'name': "Dr. Prabir Chandra Pradhan",
            'pdes': "Veterinary Officer, BAHC",
            'pest': "Sub-Divisional and Block Level Set up of Howrah",
            'pblk': "Uluberia-I",
            'pdist': "Howrah",
            'ppost': "Veterinary Officer, BAHC, Uluberia-I, Howrah",
            'psu': "Nil",
            'basis': "Transfer",
            'sub': "Block Livestock Development Officer, Sub-Divisional and Block Level Set up of Howrah, Uluberia-I, Howrah",
            'su': "Nil",
            'rem': "Transferred as Block Livestock Development Officer, Uluberia-I, Howrah",
            'deb': "bldo uluberia-I, Howrah"
        },
        {
            'rsl': "-",
            'name': "Dr. Palash Biswas",
            'pdes': "Veterinary Officer, ABAHC",
            'pest': "Sub-Divisional and Block Level Set up of Coochbehar District",
            'pblk': "Tufanganj",
            'pdist': "Coochbehar",
            'ppost': "Veterinary Officer, ABAHC, Tufanganj, Coochbehar",
            'psu': "Nil",
            'basis': "Transfer",
            'sub': "Block Livestock Development Officer, Sub-Divisional and Block Level Set up of North 24 Parganas District, Gaighata, North 24 Parganas",
            'su': "Nil",
            'rem': "Transferred as Block Livestock Development Officer, Gaighata, North 24 Parganas",
            'deb': "bldo gaighata, n24 pgs"
        },
        {
            'rsl': "-",
            'name': "Dr. Ranjit Kumar Panja",
            'pdes': "Block Livestock Development Officer",
            'pest': "Sub-Divisional and Block Level Set up of Purba Medinipur",
            'pblk': "Panskura-II",
            'pdist': "Purba Medinipur",
            'ppost': "Block Livestock Development Officer, Panskura-II, Purba Medinipur",
            'psu': "Nil",
            'basis': "Transfer",
            'sub': "Block Livestock Development Officer, Sub-Divisional and Block Level Set up of Purba Medinipur, Panskura-II, Purba Medinipur",
            'su': "Nil",
            'rem': "Transferred as Block Livestock Development Officer, Panskura-II, Purba Medinipur",
            'deb': "paskura-II bldo, paschim medinipur"
        },
        {
            'rsl': "-",
            'name': "Dr. Joydev Bera",
            'pdes': "District Veterinary Officer",
            'pest': "Office of the Joint Director, ARD, Dakshin Dinajpur",
            'pblk': "",
            'pdist': "Dakshin Dinajpur",
            'ppost': "District Veterinary Officer, Dakshin Dinajpur",
            'psu': "Nil",
            'basis': "Transfer",
            'sub': "Assistant Director, ARD (VR&I), Barasat, North 24 Parganas",
            'su': "Nil",
            'rem': "Transferred as Assistant Director, ARD (VR&I), Barasat, North 24 Parganas",
            'deb': "AD VRI Barasat"
        },
        {
            'rsl': "-",
            'name': "Dr. Biplab Kumar Maity",
            'pdes': "Veterinary Officer, SAHC",
            'pest': "Sub-Divisional and Block Level Set up of Paschim Medinipur",
            'pblk': "Chandrakona-I",
            'pdist': "Paschim Medinipur",
            'ppost': "Veterinary Officer, SAHC, Khirpai, Chandrakona-I, Paschim Medinipur",
            'psu': "Nil",
            'basis': "Transfer",
            'sub': "Deputy Director, ARD, District Office, Purba Medinipur",
            'su': "Nil",
            'rem': "Transferred as Deputy Director, ARD, District Office, Purba Medinipur",
            'deb': "DDARD Purba medinipur"
        },
        {
            'rsl': "-",
            'name': "Dr. Santanu Gharai",
            'pdes': "Veterinary Officer, BAHC",
            'pest': "Sub-Divisional and Block Level Set up of Uttar Dinajpur District",
            'pblk': "Goalpokher-II",
            'pdist': "Uttar Dinajpur",
            'ppost': "Veterinary Officer, BAHC, Asuragarh, Goalpokher-II, Uttar Dinajpur",
            'psu': "Nil",
            'basis': "Retention",
            'sub': "Veterinary Officer, BAHC, Asuragarh, Goalpokher-II, Uttar Dinajpur",
            'su': "Nil",
            'rem': "Retained at present post",
            'deb': "?"
        }
    ]

    for eo in extra_officers:
        if eo['name'].lower() not in seen_names:
            master_records.append(eo)
            seen_names.add(eo['name'].lower())

    out_r = 2
    for rec in master_records:
        out_sl = out_r - 1
        row_vals = [
            out_sl,
            rec['rsl'],
            rec['name'],
            rec['pdes'],
            rec['pest'],
            rec['pblk'],
            rec['pdist'],
            rec['ppost'],
            rec['psu'],
            rec['basis'],
            rec['sub'],
            rec['su'],
            rec['rem'],
            rec['deb']
        ]

        if "stay at present" in rec['rem'].lower():
            row_fill = fill_stay
        elif "rehabilitated" in rec['rem'].lower() or "post abolished" in rec['rem'].lower():
            row_fill = fill_rehab
        elif rec['su'] != "Nil" and rec['su'] != "":
            row_fill = fill_trans
        else:
            row_fill = fill_white

        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_master.cell(out_r, c_idx, val)
            cell.font = font_bold if c_idx in [1, 2, 3, 11] else font_data
            cell.border = thin_border
            cell.fill = row_fill
            if c_idx in [1, 2, 6, 7, 10]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        ws_master.row_dimensions[out_r].height = 24
        out_r += 1

    col_widths = {
        1: 8,   2: 14,  3: 28,  4: 25,  5: 35,  6: 18,  7: 20,
        8: 45,  9: 25, 10: 22, 11: 45, 12: 45, 13: 40, 14: 35
    }
    for c_idx, width in col_widths.items():
        ws_master.column_dimensions[get_column_letter(c_idx)].width = width

    ws_master.freeze_panes = "C2"

    # Sheet 2: 4_Column_Posting_Order
    ws_4col = wb.create_sheet(title="4_Column_Posting_Order")

    ws_4col.cell(1, 1, "GOVERNMENT OF WEST BENGAL").font = font_title
    ws_4col.cell(2, 1, "ANIMAL RESOURCES DEVELOPMENT DEPARTMENT").font = font_sub
    ws_4col.cell(3, 1, "NOTIFICATION / PROMOTION & POSTING ORDER (DEFINITIVE MASTER SCHEDULE)").font = font_sub

    order_headers = [
        "Sl. No.",
        "Name & Present Post of the Officer",
        "New Substantive Post on Promotion / Transfer",
        "Station of Service Utilization (SU) / Administrative Directives"
    ]
    for c_idx, h in enumerate(order_headers, 1):
        cell = ws_4col.cell(5, c_idx, h)
        cell.font = font_header
        cell.fill = fill_slate
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws_4col.row_dimensions[5].height = 28

    r_4 = 6
    for rec_idx, rec in enumerate(master_records, 1):
        name_val = rec['name']
        pres_post = rec['ppost']
        sub_post = rec['sub']
        su_post = rec['su']
        remarks_val = rec['rem']

        col2_text = f"{name_val}\n[Present: {pres_post}]"
        col3_text = str(sub_post)
        if su_post and su_post != "Nil":
            col4_text = f"Service Utilized at:\n{su_post}\n({remarks_val})"
        else:
            col4_text = f"{remarks_val}"

        c1 = ws_4col.cell(r_4, 1, rec_idx)
        c2 = ws_4col.cell(r_4, 2, col2_text)
        c3 = ws_4col.cell(r_4, 3, col3_text)
        c4 = ws_4col.cell(r_4, 4, col4_text)

        c1.alignment = Alignment(horizontal="center", vertical="center")
        c2.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c3.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c4.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        for c in [c1, c2, c3, c4]:
            c.font = font_data
            c.border = thin_border

        ws_4col.row_dimensions[r_4].height = 36
        r_4 += 1

    ws_4col.column_dimensions['A'].width = 10
    ws_4col.column_dimensions['B'].width = 45
    ws_4col.column_dimensions['C'].width = 45
    ws_4col.column_dimensions['D'].width = 45
    ws_4col.freeze_panes = "A6"

    # Save to all local destinations
    os.makedirs(os.path.dirname(OUTPUT_EXCEL_1), exist_ok=True)
    os.makedirs(os.path.dirname(OUTPUT_EXCEL_3), exist_ok=True)

    for target in [OUTPUT_EXCEL_1, OUTPUT_EXCEL_2, OUTPUT_EXCEL_3, OUTPUT_EXCEL_4, ALIAS_1245_1, ALIAS_1245_2]:
        wb.save(target)
        print(f"Saved: {target}")

    print(f"Master files written successfully: {len(master_records)} total officers.")

if __name__ == "__main__":
    build_master()
